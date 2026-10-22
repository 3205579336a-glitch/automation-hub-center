from __future__ import annotations

import importlib.util
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch


SCRIPT = Path(__file__).parents[1] / "resources" / "rpa" / "me01_source_list.py"
SPEC = importlib.util.spec_from_file_location("me01_source_list", SCRIPT)
assert SPEC and SPEC.loader
ME01 = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(ME01)


class FakeCell:
    def __init__(self, value=None):
        self.Value = value


class FakeCells:
    def __init__(self, values):
        self.values = values
        self.calls = 0

    def __call__(self, row, column):
        self.calls += 1
        return self.values.setdefault((row, column), FakeCell())


class FakeWorksheets:
    def __init__(self, sheet):
        self.sheet = sheet

    def __call__(self, _name_or_index):
        return self.sheet


class FakeWorkbook:
    def __init__(self):
        values = {
            (1, 1): FakeCell("Material No."),
            (1, 2): FakeCell("Parma"),
            (1, 3): FakeCell("ME01 Status"),
            (1, 4): FakeCell("Error Message"),
            (1, 5): FakeCell("Updated At"),
            (2, 1): FakeCell("16808656"),
            (2, 2): FakeCell("41889"),
            (1, 6): FakeCell("Dual Sourcing"),
            (2, 6): FakeCell("Yes"),
        }
        sheet = SimpleNamespace(
            Cells=FakeCells(values),
            UsedRange=SimpleNamespace(Columns=SimpleNamespace(Count=6), Rows=SimpleNamespace(Count=2)),
        )
        self.Worksheets = FakeWorksheets(sheet)
        self.saved_copy = None
        self.save_count = 0
        self.ReadOnly = False
        self.Close = Mock()

    def SaveCopyAs(self, path):
        self.saved_copy = path
        import openpyxl
        book = openpyxl.Workbook()
        book.active.title = 'ME01_Input'
        for (row, column), cell in self.Worksheets(1).Cells.values.items():
            book.active.cell(row, column, cell.Value)
        book.save(path)
        book.close()

    def Save(self):
        self.save_count += 1




class Me01Tests(unittest.TestCase):
    def test_hub_reuses_verified_production_window_on_another_transaction(self):
        from sap_gui_fakes import application, connection, session
        from sap_gui_connection import SapGuiConnector
        current = session(transaction='ME11')
        app = application([connection([session(system='CEQ')], 'CEQ - QA [321]'), connection([current])])
        connector = SapGuiConnector('VCE - One Digital Core [949]', 'VCE', '100')
        with patch.object(connector, 'application', return_value=app):
            self.assertIs(ME01.find_hub_session(connector), current)
        app.OpenConnection.assert_not_called()

    def test_hub_does_not_turn_session_access_errors_into_an_empty_inventory(self):
        from sap_gui_fakes import application, connection
        from sap_gui_connection import SapGuiConnector, SapGuiIssue
        app = application([connection(disabled=True)])
        connector = SapGuiConnector('VCE - One Digital Core [949]', 'VCE', '100')
        with patch.object(connector, 'application', return_value=app), self.assertRaisesRegex(SapGuiIssue, 'SAP_SCRIPTING_DISABLED'):
            ME01.find_hub_session(connector)
        app.OpenConnection.assert_not_called()

    def bulk_workbook(self, rows=201):
        workbook = FakeWorkbook()
        sheet = workbook.Worksheets(1)
        sheet.UsedRange.Rows.Count = rows
        sheet.range_calls = []
        def read(address):
            sheet.range_calls.append(address)
            match = re.fullmatch(r'A(\d+):([A-Z]+)(\d+)', address)
            first, letters, last = match.groups()
            columns = 0
            for letter in letters:
                columns = columns*26 + ord(letter)-64
            values = tuple(tuple(sheet.Cells.values.get((r,c), FakeCell()).Value
                                 for c in range(1,columns+1)) for r in range(int(first),int(last)+1))
            return SimpleNamespace(Value2=values[0][0] if int(first)==int(last) and columns==1 else values)
        sheet.Range = read
        return workbook

    def test_open_excel_bulk_read_avoids_per_cell_round_trips_for_blank_template(self):
        workbook = self.bulk_workbook()
        with tempfile.TemporaryDirectory() as directory, patch.object(ME01, 'find_open_excel_workbook', return_value=workbook):
            results = ME01.WorkbookResults(Path(directory)/'source.xlsx', Path(directory)/'backup.xlsx')
        sheet = workbook.Worksheets(1)
        self.assertEqual(sheet.range_calls, ['A1:F1','A1:F1','A2:F201'])
        self.assertEqual(sheet.Cells.calls, 0)
        self.assertEqual(results.assignments['16808656']['rows'], [2])
        self.assertTrue(results.assignments['16808656']['dualSourcing'])

    def test_bulk_read_chunks_keep_original_row_numbers_and_detect_conflicts(self):
        workbook = self.bulk_workbook(5000)
        sheet = workbook.Worksheets(1)
        sheet.Cells.values.update({(4999,1):FakeCell('20503551'),(4999,2):FakeCell('6587'),(4999,6):FakeCell('No')})
        with tempfile.TemporaryDirectory() as directory, patch.object(ME01, 'find_open_excel_workbook', return_value=workbook):
            results = ME01.WorkbookResults(Path(directory)/'source.xlsx', Path(directory)/'backup.xlsx')
            self.assertEqual(results.assignments['20503551']['rows'], [4999])
            self.assertFalse(results.assignments['20503551']['dualSourcing'])
            self.assertEqual(sheet.range_calls[-2:], ['A2:F4097','A4098:F5000'])
            self.assertEqual(sheet.Cells.calls, 0)
            sheet.Cells.values.update({(5000,1):FakeCell('20503551'),(5000,2):FakeCell('6587'),(5000,6):FakeCell('Yes')})
            with self.assertRaisesRegex(RuntimeError, 'conflicting'):
                ME01.WorkbookResults(Path(directory)/'source.xlsx', Path(directory)/'backup.xlsx')

    def test_bulk_range_checks_rectangle_shape_and_single_cell(self):
        sheet = self.bulk_workbook().Worksheets(1)
        self.assertEqual(ME01.com_range_values(sheet,1,1,1),(('Material No.',),))
        self.assertEqual(ME01.excel_column(27),'AA')
        sheet.Range = lambda _address: SimpleNamespace(Value2=((1,),))
        with self.assertRaisesRegex(RuntimeError, 'incomplete'):
            ME01.com_range_values(sheet,2,3,2)
        sheet.Range = lambda _address: (_ for _ in ()).throw(RuntimeError('Excel disconnected'))
        with self.assertRaisesRegex(RuntimeError, 'disconnected'):
            ME01.com_range_values(sheet,2,3,2)

    def test_open_excel_workbook_uses_com_and_stays_open(self):
        workbook = FakeWorkbook()
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            backup = Path(directory) / "backup.xlsx"
            with patch.object(ME01, "find_open_excel_workbook", return_value=workbook):
                results = ME01.WorkbookResults(source, backup)
            self.assertEqual(results.mode, "excel-com")
            self.assertEqual(results.assignments["16808656"]["parma"], "41889")
            results.write([2], "SUCCESS", "")
            results.save()
            self.assertEqual(workbook.Worksheets(1).Cells(2, 3).Value, "SUCCESS")
            self.assertEqual(workbook.save_count, 1)
            self.assertEqual(workbook.saved_copy, str(backup))
            workbook.Close.assert_not_called()

    def test_rot_finds_exact_workbook_in_second_excel_instance(self):
        path = Path('C:/fixtures/source.xlsx').resolve()
        target = FakeWorkbook()
        target.FullName = str(path)
        wrong = SimpleNamespace(FullName=str(path.parent / 'another.xlsx'), ReadOnly=False)
        readonly = SimpleNamespace(FullName=str(path), ReadOnly=True)
        active = SimpleNamespace(Workbooks=SimpleNamespace(Count=0))
        monikers = [SimpleNamespace(GetDisplayName=lambda *_: str(path.parent / 'another.xlsx')),
                    SimpleNamespace(GetDisplayName=lambda *_: '!' + str(path))]
        rot = SimpleNamespace(EnumRunning=lambda: monikers, GetObject=Mock(return_value=target))
        pythoncom = SimpleNamespace(GetRunningObjectTable=lambda: rot, CreateBindCtx=lambda _: object())
        with patch.dict(sys.modules, {'pythoncom': pythoncom}), \
             patch.object(ME01.win32com.client, 'GetActiveObject', return_value=active), \
             patch.object(ME01.win32com.client, 'Dispatch', side_effect=lambda obj: obj), \
             patch.object(ME01.win32com.client, 'GetObject') as open_file:
            self.assertIs(ME01.find_open_excel_workbook(path), target)
            rot.GetObject.assert_called_once_with(monikers[1])
            open_file.assert_not_called()
            active.Workbooks = Mock(Count=2, side_effect=[wrong, readonly])
            self.assertIs(ME01.find_open_excel_workbook(path), target, 'Prefer writable ROT match to readonly active match')
        target.Close.assert_not_called()

    def test_excel_busy_retry_is_bounded_and_does_not_retry_permission_errors(self):
        busy = Exception('Excel busy')
        busy.hresult = -2147418111
        action = Mock(side_effect=[busy, busy, 'saved'])
        with patch.object(ME01.time, 'sleep') as sleep:
            self.assertEqual(ME01.excel_call(action), 'saved')
        self.assertEqual(action.call_count, 3)
        self.assertEqual(sleep.call_count, 2)
        denied = Mock(side_effect=PermissionError('access denied'))
        with self.assertRaises(PermissionError): ME01.excel_call(denied)
        denied.assert_called_once()

    def test_readonly_or_failed_com_save_preserves_results_in_owned_logs(self):
        import openpyxl
        for readonly in (True, False):
            with self.subTest(readonly=readonly), tempfile.TemporaryDirectory() as folder:
                source, backup = Path(folder)/'source.xlsx', Path(folder)/'logs'/'backup.xlsx'
                backup.parent.mkdir()
                workbook = FakeWorkbook()
                workbook.ReadOnly = readonly
                with patch.object(ME01, 'find_open_excel_workbook', return_value=workbook):
                    results = ME01.WorkbookResults(source, backup)
                workbook.Save = Mock(side_effect=PermissionError('Workbook locked'))
                results.write([2], 'SKIPPED', 'SOURCE_EXISTS')
                results.save()
                self.assertEqual(results.mode, 'result-copy')
                self.assertEqual(results.result_path.parent, backup.parent)
                self.assertIn('do not rerun SAP', results.writeback_warning)
                saved = openpyxl.load_workbook(results.result_path)
                self.assertEqual(saved.active['C2'].value, 'SKIPPED')
                self.assertEqual(saved.active['D2'].value, 'SOURCE_EXISTS')
                saved.close(); results.workbook.close()
                workbook.Close.assert_not_called()

    def test_later_save_failure_retains_all_previously_completed_outputs(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as folder:
            workbook = FakeWorkbook()
            sheet = workbook.Worksheets(1)
            sheet.Cells(3, 1).Value, sheet.Cells(3, 2).Value, sheet.Cells(3, 6).Value = '20503551', '6587', 'No'
            sheet.UsedRange.Rows.Count = 3
            with patch.object(ME01, 'find_open_excel_workbook', return_value=workbook):
                results = ME01.WorkbookResults(Path(folder)/'source.xlsx', Path(folder)/'backup.xlsx')
            results.write([2], 'SUCCESS', '')
            results.save()
            workbook.Save = Mock(side_effect=PermissionError('became locked'))
            results.write([3], 'SKIPPED', 'SOURCE_EXISTS')
            results.save()
            saved = openpyxl.load_workbook(results.result_path)
            self.assertEqual([saved.active[f'C{r}'].value for r in (2, 3)], ['SUCCESS', 'SKIPPED'])
            saved.close(); results.workbook.close()
            workbook.Close.assert_not_called()

    def test_editing_input_or_headers_during_run_does_not_write_to_wrong_row(self):
        import openpyxl
        for cell, value in [((2, 2), '99999'), ((1, 3), 'User notes')]:
            with self.subTest(cell=cell), tempfile.TemporaryDirectory() as folder:
                workbook = FakeWorkbook()
                with patch.object(ME01, 'find_open_excel_workbook', return_value=workbook):
                    results = ME01.WorkbookResults(Path(folder)/'source.xlsx', Path(folder)/'backup.xlsx')
                workbook.Worksheets(1).Cells(*cell).Value = value
                results.write([2], 'SUCCESS', '')
                results.save()
                self.assertEqual(results.mode, 'result-copy')
                self.assertIsNone(workbook.Worksheets(1).Cells(2, 3).Value)
                self.assertEqual(workbook.save_count, 0)
                saved = openpyxl.load_workbook(results.result_path)
                self.assertEqual(saved.active['B2'].value, '41889')
                self.assertEqual(saved.active['C2'].value, 'SUCCESS')
                saved.close(); results.workbook.close()
                workbook.Close.assert_not_called()

    def test_workbook_opened_mid_run_attaches_without_closing_excel(self):
        import openpyxl
        for changed in (False, True):
            with self.subTest(changed=changed), tempfile.TemporaryDirectory() as folder:
                source, backup = Path(folder)/'source.xlsx', Path(folder)/'backup.xlsx'
                live = FakeWorkbook()
                live.SaveCopyAs(str(source))
                original = source.read_bytes()
                with patch.object(ME01, 'find_open_excel_workbook', return_value=None):
                    results = ME01.WorkbookResults(source, backup)
                live.Worksheets(1).Cells(2, 2).Value = '99999' if changed else '41889'
                live.Worksheets(1).Cells(3, 9).Value = 'unsaved user notes'
                results.write([2], 'SUCCESS', '')
                with patch.object(ME01, 'find_open_excel_workbook', return_value=live):
                    results.save()
                self.assertEqual(source.read_bytes(), original, 'No direct overwrite of the now-open source file')
                live.Close.assert_not_called()
                self.assertEqual(live.Worksheets(1).Cells(3, 9).Value, 'unsaved user notes')
                if changed:
                    self.assertEqual(results.mode, 'result-copy')
                    self.assertEqual(live.save_count, 0)
                    self.assertIsNone(live.Worksheets(1).Cells(2, 3).Value)
                    results.workbook.close()
                else:
                    self.assertEqual(results.mode, 'excel-com')
                    self.assertEqual(live.save_count, 1)
                    self.assertEqual(live.Worksheets(1).Cells(2, 3).Value, 'SUCCESS')

    def test_locked_file_save_uses_result_copy_without_overwriting_source(self):
        with tempfile.TemporaryDirectory() as folder:
            source, backup = Path(folder)/'source.xlsx', Path(folder)/'backup.xlsx'
            FakeWorkbook().SaveCopyAs(str(source))
            original = source.read_bytes()
            with patch.object(ME01, 'find_open_excel_workbook', return_value=None):
                results = ME01.WorkbookResults(source, backup)
                results.write([2], 'SUCCESS', '')
                with patch.object(results.workbook, 'save', side_effect=PermissionError('locked by Excel')):
                    results.save()
            self.assertEqual(source.read_bytes(), original)
            self.assertEqual(results.mode, 'result-copy')
            results.write([2], 'SKIPPED', 'SOURCE_EXISTS')
            with patch.object(results.workbook, 'save', side_effect=PermissionError('log folder unwritable')):
                with self.assertRaisesRegex(PermissionError, 'log folder'): results.save()
            results.workbook.close()

    def test_existing_target_is_skipped_and_next_material_is_processed(self):
        from collections import OrderedDict
        import hub_artifacts
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder)/'source.xlsx'
            source.touch()
            results = Mock(assignments=OrderedDict([
                ('16808656', {'material': '16808656', 'parma': '41889', 'dualSourcing': False, 'rows': [2]}),
                ('20503551', {'material': '20503551', 'parma': '6587', 'dualSourcing': True, 'rows': [3]})]),
                result_path=source, mode='excel-com', writeback_warning='')
            with patch.dict(os.environ, {'EXCEL_PATH': str(source), 'HUB_RUN_ID': ''}), \
                 patch.object(hub_artifacts, 'artifact_directory', return_value=Path(folder)), \
                 patch.object(ME01, 'WorkbookResults', return_value=results), \
                 patch.object(ME01, 'find_active_session', return_value=object()), \
                 patch.object(ME01, 'open_me01'), patch.object(ME01, 'emit') as emit, \
                 patch.object(ME01, 'maintain_material', side_effect=[False, True]) as maintain:
                self.assertEqual(ME01.main(), 0)
            self.assertEqual(maintain.call_count, 2)
            self.assertEqual(results.write.call_args_list[0].args[:2], ([2], 'SKIPPED'))
            self.assertEqual(results.write.call_args_list[1].args[:2], ([3], 'SUCCESS'))
            complete = [call.kwargs for call in emit.call_args_list if call.args[0] == 'complete'][0]
            self.assertEqual((complete['processed'], complete['skipped'], complete['succeeded'], complete['failed']), (2, 1, 1, 0))






    def test_dual_sourcing_is_required_and_duplicate_modes_must_match(self):
        for value in [None, '', 'maybe']:
            with self.assertRaises(RuntimeError): ME01.parse_dual_sourcing(value, 2)
        self.assertFalse(ME01.parse_dual_sourcing(False, 2))
        self.assertTrue(ME01.parse_dual_sourcing('✓', 2))
        from collections import OrderedDict
        results = object.__new__(ME01.WorkbookResults)
        results.assignments = OrderedDict()
        results.input_rows = {}
        results._add_assignment('16808656', '41889', 2, 'Yes')
        with self.assertRaisesRegex(RuntimeError, 'conflicting'):
            results._add_assignment('16808656', '41889', 3, 'No')



    def test_openpyxl_mode_reads_required_choice_and_rejects_old_template(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as folder:
            source, backup = Path(folder) / 'input.xlsx', Path(folder) / 'backup.xlsx'
            w = openpyxl.Workbook()
            w.active.title = 'ME01_Input'
            w.active.append(['Material No.', 'Parma', 'Dual Sourcing', 'ME01 Status'])
            w.active.append(['16808656', '41889', 'No', 'SUCCESS'])
            w.save(source)
            with patch.object(ME01, 'find_open_excel_workbook', return_value=None):
                results = ME01.WorkbookResults(source, backup)
                self.assertFalse(results.assignments['16808656']['dualSourcing'])
                results.workbook.close()
                w.active.cell(2,3).value = ''
                w.save(source)
                with self.assertRaisesRegex(RuntimeError, 'required'):
                    ME01.WorkbookResults(source, backup)
            w.close()

    def test_download_template_preserves_table_annotations_and_required_validation(self):
        import openpyxl
        w = openpyxl.load_workbook(SCRIPT.parent.parent / 'templates' / 'ME01_Source_List_Template.xlsx')
        try:
            s = w['ME01_Input']
            self.assertEqual([c.value for c in s[1]], ['Material No.', 'Parma', 'Dual Sourcing', 'ME01 Status', 'Error Message', 'Updated At'])
            self.assertEqual(s.tables['Me01SourceListInput'].ref, 'A1:F201')
            self.assertIsNone(s['C1'].comment)
            self.assertTrue(all(s[c].comment is not None for c in ['A1','B1','D1','E1','F1']))
            v = list(s.data_validations.dataValidation)[0]
            self.assertEqual(str(v.sqref), 'C2:C1001')
            self.assertEqual(v.formula1, '"Yes,No"')
            self.assertTrue(v.showErrorMessage)
            self.assertFalse(v.allow_blank)
        finally:
            w.close()




    def test_sap_logon_is_launched_when_no_session_exists(self):
        expected_session = object()
        with patch.object(ME01, "try_active_session", side_effect=[None, expected_session]), \
             patch.object(ME01, "launch_sap_logon") as launch:
            actual_session = ME01.find_active_session()
        launch.assert_called_once_with()
        self.assertIs(actual_session, expected_session)


if __name__ == "__main__":
    unittest.main()
