"""Offline RFQ validation/adapter tests. SAP is always replaced by a mock."""
import ast
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch, MagicMock

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'resources' / 'rpa'))
sys.argv = [sys.argv[0], '--hub', '--validate-only']
import rfq_engine as engine
from rfq_runtime import initialize, run_hub, validate_workbook


class RfqTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / 'input.xlsx'
        self.args = initialize(Path(self.temp.name) / '.env')
        self.args.excel_path = str(self.path)
        self.args.run_id = 'offline-test'
        self.args.stop_file = str(Path(self.temp.name) / 'stop')

    def tearDown(self):
        self.temp.cleanup()

    def book(self, rows=None, headers=None, sheet='RPA_Input'):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet
        columns = list(headers or ['Plant', 'Project No.', 'Material', 'Supplier Parma', 'Supplier Email', 'Quotation Due Date'])
        defaults = {'PPAP Date': '2030-12-31', 'Technology': 'TEST_TECH', '12 MR Qty': 120, 'RFQ Qty Serial': 25}
        extra = [name for name in defaults if name not in columns]
        ws.append(columns + extra)
        for row in rows or [['C100', '775773', '12345678', '26517', 'buyer@example.com', '2030-12-31']]:
            ws.append(list(row) + [None] * (len(columns) - len(row)) + [defaults[name] for name in extra])
        wb.save(self.path)
        wb.close()

    def preview(self):
        return validate_workbook(engine, self.path)

    def test_no_env_no_sap_no_source_writes(self):
        self.book()
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        with patch.object(engine, 'SapSession', side_effect=AssertionError('SAP must not be opened')):
            result = self.preview()
        self.assertEqual((result['validRows'], result['groupCount']), (1, 1))
        self.assertEqual(before, hashlib.sha256(self.path.read_bytes()).hexdigest())
        self.assertEqual(list(Path(self.temp.name).iterdir()), [self.path])

    def test_template_is_blank_and_compatible(self):
        preview = validate_workbook(engine, ROOT / 'resources/templates/Create_RFQ_Template.xlsx')
        self.assertEqual(preview['totalRows'], 0)

    def test_supplied_sap_business_definitions_unchanged(self):
        # Baseline AST digest calculated from the user's original v32 source.
        # Exclude startup/main, ExcelStore and three authorized checkbox/Qty
        # edits (No clears, Yes preserves, optional blank Qty is not written).
        tree = ast.parse((ROOT/'resources/rpa/rfq_engine.py').read_text(encoding='utf-8-sig').replace('v31', 'v32'))
        # Authorized query reset fix: exclude only these SapSession methods,
        # keeping every other SAP/session business method baseline-protected.
        query_methods = {'set_material_multiple_selection', '_execute_npl_query', '_replace_material_multiple_selection'}
        for node in tree.body:
            if isinstance(node, ast.ClassDef) and node.name == 'SapSession':
                node.body = [method for method in node.body
                             if not isinstance(method, (ast.FunctionDef, ast.AsyncFunctionDef)) or method.name not in query_methods]
        definitions = [ast.dump(n, include_attributes=False) for n in tree.body
                       if isinstance(n, (ast.FunctionDef, ast.ClassDef)) and n.name not in {'main', 'standalone_main', 'ExcelStore', 'prepare_prod_direct_rfq_rows', 'fill_rfq_grid_rows', 'fill_buyer_receipt_rows'}]
        self.assertEqual(len(definitions), 68)
        self.assertEqual(hashlib.sha256('\n'.join(definitions).encode()).hexdigest(),
                         'e86af6b0646d42af572f74da55a1b0330cf101212f9377ad78218f20c295faa7')

    def test_environment_guard_rejects_wrong_system_client_or_confirmation(self):
        session = engine.SapSession.__new__(engine.SapSession)
        session.connection_target_verified = True
        session.connection_description = 'Mock VCE'
        session.system_name, session.client, session.user = 'VCE', '100', 'TEST_ONLY'
        with patch.object(engine, 'ALLOW_PRODUCTION_WRITE', True):
            session._validate_environment()
            for attribute, invalid in [('system_name', 'CEQ'), ('client', '200'), ('connection_target_verified', False)]:
                with patch.object(session, attribute, invalid), self.assertRaises(RuntimeError):
                    session._validate_environment()
        with patch.object(engine, 'ALLOW_PRODUCTION_WRITE', False), self.assertRaises(RuntimeError):
            session._validate_environment()

    def test_non_green_buyer_receipt_is_not_accepted(self):
        self.book()
        self.preview()
        store = engine.ExcelStore(self.path, 'RPA_Input')
        try:
            tasks = store.load_tasks()
            tasks[0].buyer_receipt_row = 0
            sap = MagicMock()
            sap.exists.return_value = False
            sap.map_material_rows.return_value = ({tasks[0].material: [0]}, {})
            with patch.object(engine, 'green_status_debug', return_value=(False, 'not green')), \
                 patch.object(engine, 'BUYER_GREEN_STATUS_TIMEOUT_SEC', 0):
                accepted, _, _ = engine.save_buyer_receipt(sap, object(), tasks)
            self.assertEqual(accepted, [])
            sap.wait_and_confirm_buyer_receipt_save.assert_called_once()
        finally:
            store.workbook.close()

    def test_group_51_materials_split_50_1(self):
        self.book([['C100', '775773', str(10000000+i), '26517', 'buyer@example.com', '2030-12-31'] for i in range(51)])
        result = self.preview()
        self.assertEqual([len(g['materials']) for g in result['groups']], [50, 1])

    def test_three_part_group_key(self):
        self.book([['C100', 'P1', '111', '26517', 'a@example.com', '2030-12-31'],
                   ['C100', 'P1', '222', '26517', 'a@example.com', '2030-12-31'],
                   ['C100', 'P2', '333', '26517', 'a@example.com', '2030-12-31'],
                   ['C200', 'P1', '444', '26517', 'a@example.com', '2030-12-31'],
                   ['C100', 'P1', '555', '34638', 'a@example.com', '2030-12-31']])
        self.assertEqual(self.preview()['groupCount'], 4)

    def test_conflicting_headers_warn_without_splitting(self):
        self.book([['C100', 'P1', '111', '26517', 'a@example.com', '2030-12-31'],
                   ['C100', 'P1', '222', '26517', 'b@example.com', '2031-01-01']])
        result = self.preview()
        self.assertEqual(result['groupCount'], 1)
        self.assertEqual(len(result['warnings']), 1)

    def test_reordered_columns_and_sheet_alias(self):
        self.book([['buyer@example.com', '26517', '12345678', '775773', '2030-12-31', 'C100']],
                  ['Supplier Email', 'Supplier Parma', 'Part Number', 'MPP Project No.', 'RFQ Due Date', 'Target Plant'], 'My Input')
        result = self.preview()
        self.assertEqual(result['validRows'], 1)
        self.assertEqual(result['sheetName'], 'My Input')
        self.assertEqual(engine.SUPPLIER_EMAIL_EXCEL_COL, 1)
        self.assertEqual(engine.RFQ_DUE_DATE_EXCEL_COL, 5)

    def test_invalid_fields_and_missing_material(self):
        cases = [(0, 'C'), (1, ''), (2, ''), (3, ''), (4, 'bad@email'), (5, 'not-a-date')]
        for column, value in cases:
            row = ['C100', 'P1', '111', '26517', 'a@example.com', '2030-12-31']
            row[column] = value
            self.book([row])
            self.assertEqual(self.preview()['invalidRows'], 1, (column, value))

    def test_duplicate_and_missing_headers(self):
        self.book(headers=['Plant', 'Project No.', 'Material', 'Intended Supplier', 'Supplier Email', 'Supplier Email'])
        with self.assertRaises(ValueError):
            self.preview()

    def test_exact_duplicate_material_is_skipped(self):
        row = ['C100', 'P1', '111', '26517', 'a@example.com', '2030-12-31']
        self.book([row, row])
        preview = self.preview()
        self.assertEqual((preview['totalRows'], preview['validRows'], preview['invalidRows'], preview['duplicateRows']), (2, 1, 0, 1))
        self.assertEqual(preview['sample'][1]['duplicateOf'], 2)

    def test_optional_business_validation(self):
        for header, value in [('Attach File', 'Yes'), ('AFM Request', 'maybe'), ('RFQ Comment', 'x'*201), ('PPAP Date', 'invalid')]:
            self.book([['C100', 'P1', '111', '26517', 'a@example.com', '2030-12-31', value]],
                      ['Plant', 'Project No.', 'Material', 'Intended Supplier', 'Supplier Email', 'Quotation Due Date', header])
            self.assertEqual(self.preview()['invalidRows'], 1, header)

    def test_excel_serial_and_text_dates(self):
        for date in ['31.12.2030', '2030-12-31', 47848]:
            self.book([['C100', 'P1', '111', '26517', 'a@example.com', date]])
            self.assertEqual(self.preview()['validRows'], 1)

    def test_success_rows_skipped(self):
        self.book([['C100', 'P1', '111', '26517', 'a@example.com', '2030-12-31', 'SUCCESS']],
                  ['Plant', 'Project No.', 'Material', 'Intended Supplier', 'Supplier Email', 'Quotation Due Date', 'RFQ Status'])
        self.assertEqual(self.preview()['skippedRows'], 1)

    def test_stale_environment_cannot_override_hub(self):
        with patch.dict(os.environ, {'SAP_TARGET_ENV': 'QA', 'EXPECTED_SAP_SYSTEM': 'BAD', 'EXPECTED_SAP_USER': 'someone',
                                    'SAP_ENVIRONMENT_GUARD': 'false', 'GROUP_RFQ_BY_PARMA': 'false', 'MAX_MATERIALS_PER_GROUP': '1'}):
            initialize(Path(self.temp.name) / '.env')
            self.assertEqual(os.environ['SAP_TARGET_ENV'], 'PROD')
            self.assertEqual(os.environ['EXPECTED_SAP_SYSTEM'], 'VCE')
            self.assertEqual(os.environ['EXPECTED_SAP_USER'], '')
            self.assertEqual(os.environ['GROUP_RFQ_BY_PARMA'], 'true')
            self.assertEqual(os.environ['MAX_MATERIALS_PER_GROUP'], '50')
            self.assertEqual(os.environ['ALLOW_PRODUCTION_WRITE'], 'false')

    def test_confirmation_and_validation_prevent_sap(self):
        self.book()
        self.args.validate_only = False
        with patch.object(engine, 'SapSession', side_effect=AssertionError('must not reach SAP')):
            self.assertEqual(run_hub(engine, self.args), 1)

    def mock_run(self, status='SUCCESS', stop=False, broken=False, partial=False):
        self.book([['C100', 'P1', '111', '26517', 'a@example.com', '2030-12-31'],
                   ['C100', 'P1', '222', '26517', 'a@example.com', '2030-12-31']])
        self.args.validate_only = False
        self.args.production_confirmed = True
        events = []
        def process(sap, excel, group):
            if broken:
                excel.write_status([t.excel_row for t in group.tasks], buyer_status='SUCCESS')
                excel.save()
                raise engine.SapRpaError('RFQ_FINAL_SAVE', 'mock unresolved popup')
            for task in group.tasks:
                needs_ppap = status == 'INPUT_REQUIRED_PPAP_DATE' or (partial and task.excel_row == 3)
                excel.write_status([task.excel_row], buyer_status='INPUT_REQUIRED_PPAP_DATE' if needs_ppap else 'SUCCESS',
                                   rfq_status='NOT_STARTED' if needs_ppap else status,
                                   rfq_number='1000123456' if status == 'SUCCESS' and not needs_ppap else '',
                                   error_message='PPAP Date needed' if needs_ppap else '')
            excel.save()
            if stop:
                Path(self.args.stop_file).touch()
            return ('SUCCESS' if status == 'SUCCESS' else 'SKIPPED'), ('1000123456' if status == 'SUCCESS' else '')
        with patch.object(engine, 'EXCEL_PATH', self.path), patch.object(engine, 'SapSession', return_value=object()), \
             patch.object(engine, 'process_group', side_effect=process), \
             patch('rfq_runtime.emit', side_effect=lambda args, kind, **data: events.append({'type': kind, **data})):
            code = run_hub(engine, self.args)
        return code, events

    def test_mock_success_events_and_write_all_group_rows(self):
        code, events = self.mock_run()
        self.assertEqual(code, 0)
        self.assertEqual(events[-1]['type'], 'RUN_COMPLETED')
        self.assertEqual(events[-1]['succeeded'], 1)
        self.assertEqual(events[-1]['rfqNumbers'], ['1000123456'])
        wb = openpyxl.load_workbook(self.path)
        ws = wb.active
        col = [c.value for c in ws[1]].index('RFQ Number') + 1
        self.assertEqual([ws.cell(r, col).value for r in (2, 3)], ['1000123456']*2)
        wb.close()

    def test_action_required_is_nonfatal(self):
        code, events = self.mock_run('INPUT_REQUIRED_PPAP_DATE')
        self.assertEqual(code, 0)
        self.assertIn('ACTION_REQUIRED', [e['type'] for e in events])
        self.assertEqual(events[-1]['skipped'], 1)
        self.assertEqual(next(e for e in events if e['type'] == 'ACTION_REQUIRED')['materials'], ['111'])

    def test_group_with_success_and_ppap_skip(self):
        code, events = self.mock_run(partial=True)
        self.assertEqual(code, 0)
        self.assertEqual(events[-1]['succeeded'], 1)
        self.assertEqual(events[-1]['withSkips'], 1)
        self.assertEqual(events[-1]['rfqNumbers'], ['1000123456'])

    def test_cancel_keeps_saved_rfq_numbers(self):
        code, events = self.mock_run(stop=True)
        self.assertEqual(code, 3)
        self.assertEqual(events[-1]['type'], 'RUN_CANCELLED')
        self.assertEqual(events[-1]['rfqNumbers'], ['1000123456'])

    def test_buyer_saved_error_keeps_group_failure(self):
        code, events = self.mock_run(broken=True)
        self.assertEqual(events[-1]['failed'], 1)
        self.assertEqual(events[-1]['succeeded'], 0)

    def test_locked_excel_fallback(self):
        self.book()
        self.preview()  # Resolve header columns.
        store = engine.ExcelStore(self.path, 'RPA_Input')
        try:
            original = store.workbook.save
            def locked(path):
                if Path(path) == self.path:
                    raise PermissionError('mock Excel lock')
                return original(path)
            with patch.object(store.workbook, 'save', side_effect=locked), patch.object(engine, 'EXCEL_SAVE_RETRIES', 2), \
                 patch.object(store, '_sync_status_to_open_excel', return_value=False):
                store.write_status([2], rfq_status='SUCCESS', rfq_number='1000123456')
                store.save()
                self.assertTrue(store.used_lock_fallback)
                self.assertTrue(store.output_path.exists())
        finally:
            store.workbook.close()

    def test_cli_utf8_no_env(self):
        self.book()
        result = subprocess.run([sys.executable, str(ROOT/'resources/rpa/automation_engine_launcher.py'), 'rfq', '--hub', '--validate-only', '--excel-path', str(self.path)], capture_output=True, encoding='utf-8', timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('HUB_EVENT:', result.stdout)


if __name__ == '__main__':
    unittest.main(argv=[sys.argv[0]], verbosity=2)
