from __future__ import annotations

import importlib.util
import json
import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "resources" / "rpa" / "me52n_project_ref.py"
spec = importlib.util.spec_from_file_location("me52n_project_ref_test", SCRIPT)
assert spec and spec.loader
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class FakeLocator:
    @property
    def first(self):
        return self

    def click(self, timeout=0):
        return None


class FakeFrame:
    def get_by_role(self, *args, **kwargs):
        return FakeLocator()

    def locator(self, *args, **kwargs):
        return FakeLocator()


class FakeKeyboard:
    def press(self, *args, **kwargs):
        return None


class FakePage:
    def __init__(self):
        self.frames = [FakeFrame()]
        self.keyboard = FakeKeyboard()


class FakeCheckpoint:
    def __init__(self, completed=()):
        self.completed = set(completed)

    def get_record(self, pr):
        return {"Status": "SUCCESS"} if pr in self.completed else None


class FakeComCell:
    def __init__(self, sheet, row, column):
        self.sheet = sheet
        self.row = row
        self.Column = column

    @property
    def Value(self):
        return self.sheet.values.get((self.row, self.Column))

    @Value.setter
    def Value(self, value):
        self.sheet.values[(self.row, self.Column)] = value

    def End(self, _direction):
        columns = [column for (row, column), value in self.sheet.values.items() if row == 1 and value]
        return FakeComCell(self.sheet, 1, max(columns, default=1))


class FakeComSheet:
    def __init__(self):
        self.values = {(1, 1): "PR", (2, 1): "1001"}
        self.Columns = type("Columns", (), {"Count": 16384})()
        self.UsedRange = type("UsedRange", (), {"Columns": type("Columns", (), {"Count": 1})()})()

    def Cells(self, row, column):
        return FakeComCell(self, row, column)


class FakeComWorkbook:
    def __init__(self, path):
        self.FullName = str(path)
        self.sheet = FakeComSheet()
        self.save_count = 0

    def Worksheets(self, _name_or_index):
        return self.sheet

    def SaveCopyAs(self, destination):
        shutil.copy2(self.FullName, destination)

    def Save(self):
        self.save_count += 1


class Me52nSafetyTests(unittest.TestCase):
    def test_event_paths_are_json_serializable(self):
        message = module.format_me52n_event({
            "status": "COMPLETE",
            "logPath": Path(r"C:\RPA Logs\result.csv"),
            "backupPath": Path(r"C:\RPA Logs\backup.xlsx"),
        })
        self.assertTrue(message.startswith("ME52N_EVENT "))
        payload = json.loads(message.removeprefix("ME52N_EVENT "))
        self.assertEqual(payload["logPath"], r"C:\RPA Logs\result.csv")
        self.assertEqual(payload["backupPath"], r"C:\RPA Logs\backup.xlsx")

    def test_worker_count_supports_real_parallelism(self):
        self.assertEqual(module.calculate_worker_count(8, 3, False), 3)
        self.assertEqual(module.calculate_worker_count(2, 5, False), 2)
        self.assertEqual(module.calculate_worker_count(8, 3, True), 1)

    def test_pr_rows_are_deduplicated_and_checkpointed(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "prs.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["PR"])
            for value in ["1001", "1001", "1002", "bad", None]:
                sheet.append([value])
            workbook.save(path)

            names = ["EXCEL_PATH", "PR_COLUMN", "DATA_START_ROW", "BATCH_START_ROW", "BATCH_SIZE", "PR_ROW_MAPPING_SCOPE"]
            old = {name: getattr(module, name) for name in names}
            try:
                module.EXCEL_PATH = str(path)
                module.PR_COLUMN = 1
                module.DATA_START_ROW = 2
                module.BATCH_START_ROW = 0
                module.BATCH_SIZE = 0
                module.PR_ROW_MAPPING_SCOPE = "CURRENT_BATCH"
                tasks, stats, skipped = module.load_pr_tasks(FakeCheckpoint({"1002"}))
            finally:
                for name, value in old.items():
                    setattr(module, name, value)

            self.assertEqual([task["pr"] for task in tasks], ["1001"])
            self.assertEqual(tasks[0]["source_rows"], [2, 3])
            self.assertEqual([item["pr"] for item in skipped], ["1002"])
            self.assertEqual(stats["duplicate_rows"], 1)
            self.assertEqual(stats["invalid_rows"], 1)

    def test_save_requires_positive_sap_confirmation(self):
        originals = module.visible, module.dismiss_continue_buttons, module.frame_text, module.SAVE_VERIFY_TIMEOUT_SEC
        try:
            module.visible = lambda locator: True
            module.dismiss_continue_buttons = lambda page, timeout_sec=0: None
            module.SAVE_VERIFY_TIMEOUT_SEC = 0.01
            module.frame_text = lambda frame: "Purchase requisition 1234567890 changed"
            self.assertTrue(module.save_purchase_requisition(FakePage())[0])

            module.frame_text = lambda frame: "Document has not been changed"
            ok, message = module.save_purchase_requisition(FakePage())
            self.assertFalse(ok)
            self.assertIn("未确认保存", message)

            module.frame_text = lambda frame: "Ready"
            ok, message = module.save_purchase_requisition(FakePage())
            self.assertFalse(ok)
            self.assertIn("未检测到", message)
        finally:
            module.visible, module.dismiss_continue_buttons, module.frame_text, module.SAVE_VERIFY_TIMEOUT_SEC = originals

    def test_open_excel_remains_open_and_is_written_through_com(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "open-prs.xlsx"
            workbook = openpyxl.Workbook()
            workbook.active.append(["PR"])
            workbook.active.append(["1001"])
            workbook.save(path)

            fake_workbook = FakeComWorkbook(path)
            names = [
                "EXCEL_PATH", "SOURCE_SHEET_NAME", "WRITEBACK_TO_SOURCE_EXCEL",
                "CREATE_SOURCE_BACKUP", "find_open_excel_workbook",
            ]
            old = {name: getattr(module, name) for name in names}
            try:
                module.EXCEL_PATH = str(path)
                module.SOURCE_SHEET_NAME = ""
                module.WRITEBACK_TO_SOURCE_EXCEL = True
                module.CREATE_SOURCE_BACKUP = True
                module.find_open_excel_workbook = lambda _path: fake_workbook
                writer = module.SourceStatusWriter()
                self.assertEqual(writer.mode, "excel-com")
                self.assertTrue(writer.backup_path and writer.backup_path.exists())
                writer.update_rows([2], "1001", "SUCCESS", 2, "Saved")
            finally:
                for name, value in old.items():
                    setattr(module, name, value)

            status_column = writer.columns[module.SOURCE_STATUS_HEADER]
            self.assertEqual(fake_workbook.sheet.Cells(2, status_column).Value, "SUCCESS")
            self.assertGreaterEqual(fake_workbook.save_count, 2)


if __name__ == "__main__":
    unittest.main()
