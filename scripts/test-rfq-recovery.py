"""Offline acceptance for exact input deduplication and checkpoint recovery."""
import copy
import hashlib
import io
import importlib.util
import json
import sys
import unittest
from contextlib import redirect_stdout, redirect_stderr
from dataclasses import fields
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl
spec = importlib.util.spec_from_file_location('rfq_integration_tests', Path(__file__).with_name('test-rfq-integration.py'))
fixture_module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixture_module)
RfqTests, engine = fixture_module.RfqTests, fixture_module.engine
from rfq_runtime import run_hub
from rfq_interaction import business_key, verify_staging, verify_saved_buyer, check_session, session_identity


class RecoveryTests(unittest.TestCase):
    def setUp(self):
        self.fixture = RfqTests()
        self.fixture.setUp()
        self.args = self.fixture.args
        self.args.validate_only = False
        self.args.production_confirmed = True
        self.args.control_dir = str(Path(self.fixture.temp.name) / 'interaction')
        self.events = []
        self.sap = MagicMock()
        self.sap.session.Id = '/app/con[0]/ses[0]'
        self.sap.session.Info.SystemName = 'VCE'
        self.sap.session.Info.Client = '100'
        self.sap.session.Info.User = 'MOCK_USER'
        self.sap.session.Busy = False
        self.sap.exists.return_value = False
        self.sap.status_bar.return_value = ('S', '')
        self.sap.column_exists.return_value = True
        self.grid = MagicMock()
        self.grid.GetCellType.return_value = 'Checkbox'
        self.grid.GetCellChangeable.return_value = True
        self.title = 'RFQ/Quotation Creation by vendor'
        self.green = True
        self.sap.find.side_effect = lambda name, **kw: type('Window', (), {'Text': self.title})() if name == 'wnd[0]' else self.grid
        values = {'WERKS': 'C100', 'MPPPSPID': 'P1', 'LIFNR1': '26517', 'Z12MRQTY': '120', 'RFQQTY_S': '25'}
        self.sap.get_cell.side_effect = lambda grid, row, column: ('@08@' if self.green else '') if column == 'ICON' else values.get(column, '')

    def tearDown(self):
        self.fixture.tearDown()

    def book(self, rows=None, headers=None):
        row = ['C100', 'P1', '111', '26517', 'a@example.com', '2030-12-31']
        self.fixture.book(rows or [row], headers)

    def reply(self, data, action='continue'):
        directory = Path(self.args.control_dir)
        target = directory / (data['requestId'] + '.json')
        target.write_text(json.dumps({'requestId': data['requestId'], 'runId': self.args.run_id, 'action': action}), encoding='utf-8')

    def invoke(self, process, on_event=None, save=None):
        def emit(args, kind, **data):
            self.events.append({'type': kind, **data})
            if on_event:
                on_event(kind, data)
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()), \
             patch.object(engine, 'EXCEL_PATH', self.fixture.path), \
             patch.object(engine, 'SapSession', return_value=self.sap), \
             patch.object(engine, 'process_group', side_effect=process), \
             patch('rfq_runtime.emit', side_effect=emit):
            if save is not None:
                with patch.object(engine, 'save_buyer_receipt', side_effect=save):
                    return run_hub(engine, self.args)
            return run_hub(engine, self.args)

    def statuses(self):
        wb = openpyxl.load_workbook(self.fixture.path)
        try:
            ws = wb.active
            headers = {c.value: c.column for c in ws[1]}
            return [{name: ws.cell(row, headers[name]).value for name in ('RFQ Status', 'RFQ Number', 'Error Stage', 'Error Message')}
                    for row in range(2, ws.max_row + 1)]
        finally:
            wb.close()

    def test_normalized_exact_duplicates_only_execute_first(self):
        self.book([['C100','P1','000111','026517','a@example.com','31.12.2030'],
                   ['C100','P1','111','26517','a@example.com','2030-12-31']])
        before = hashlib.sha256(self.fixture.path.read_bytes()).hexdigest()
        preview = self.fixture.preview()
        self.assertEqual((preview['validRows'], preview['duplicateRows'], preview['invalidRows']), (1,1,0))
        self.assertEqual(hashlib.sha256(self.fixture.path.read_bytes()).hexdigest(), before)
        calls = []
        def process(sap, excel, group):
            calls.append([t.excel_row for t in group.tasks])
            excel.write_status([t.excel_row for t in group.tasks], rfq_status='SUCCESS', rfq_number='MOCK-1')
            excel.save()
            return 'SUCCESS','MOCK-1'
        self.assertEqual(self.invoke(process), 0)
        self.assertEqual(calls, [[2]])
        result = self.statuses()
        self.assertEqual(result[1]['RFQ Status'], 'DUPLICATE_SKIPPED')
        self.assertEqual(result[1]['Error Stage'], 'INPUT_DEDUPLICATION')
        self.assertEqual(result[1]['Error Message'], 'Duplicate of Excel row 2')
        self.assertFalse(any(e['type'] == 'ACTION_REQUIRED' for e in self.events))

    def test_same_material_with_changed_inputs_is_not_deduplicated(self):
        headers = ['Plant','Project No.','Material','Supplier Parma','Supplier Email','Quotation Due Date','PPAP Date','Technology','12 MR Qty','RFQ Qty Prototype','RFQ Qty Serial','Cost Breakdown','RFQ Comment','AFM Request']
        original = ['C100','P1','111','26517','a@example.com','2030-12-31','2030-12-31','TECH',120,0,25,'No','original','No']
        for col, value in [(4,'b@example.com'),(5,'2030-12-30'),(6,'2030-12-30'),(7,'OTHER'),(8,121),(9,1),(10,26),(11,'Yes'),(12,'changed'),(13,'Yes')]:
            with self.subTest(column=headers[col]):
                other = list(original)
                other[col] = value
                self.book([original, other], headers)
                preview = self.fixture.preview()
                self.assertEqual((preview['duplicateRows'],preview['invalidRows']), (0,1))
                self.assertIn('Different business data', preview['sample'][1]['message'])

    def test_all_optional_task_inputs_are_in_fingerprint(self):
        self.book()
        self.fixture.preview()
        store = engine.ExcelStore(self.fixture.path, 'RPA_Input')
        try:
            task = store.load_tasks()[0]
        finally:
            store.close()
        metadata = {'excel_row','npl_row','buyer_receipt_row','rfq_row'}
        for field in fields(task):
            if field.name in metadata:
                continue
            changed = copy.deepcopy(task)
            value = getattr(changed, field.name)
            setattr(changed, field.name, not value if isinstance(value,bool) else [*value,'different'] if isinstance(value,list) else str(value)+'different')
            self.assertNotEqual(business_key(task), business_key(changed),field.name)

    def test_copy_of_completed_row_is_not_executed(self):
        headers = ['Plant','Project No.','Material','Supplier Parma','Supplier Email','Quotation Due Date','RFQ Status']
        base = ['C100','P1','111','26517','a@example.com','2030-12-31']
        self.book([base+[''], base+['SUCCESS']],headers)
        preview = self.fixture.preview()
        self.assertEqual(preview['validRows'],0)
        self.assertEqual(preview['duplicateRows'],1)

    def test_quantity_equivalence_preserves_blank_vs_zero(self):
        self.book()
        self.fixture.preview()
        store = engine.ExcelStore(self.fixture.path,'RPA_Input')
        try: task = store.load_tasks()[0]
        finally: store.close()
        same = copy.deepcopy(task)
        same.qty_12mr = '120.000'
        self.assertEqual(business_key(task),business_key(same))
        changed = copy.deepcopy(task)
        changed.rfq_qty_p = '0'
        self.assertNotEqual(business_key(task),business_key(changed))
        task.qty_12mr = '123456789012345678901234567891'
        changed.qty_12mr = '123456789012345678901234567892'
        changed.rfq_qty_p = task.rfq_qty_p
        self.assertNotEqual(business_key(task),business_key(changed))

    def test_staging_failed_check_then_resume_without_duplicate_create(self):
        self.book()
        duplicate = True
        self.sap.map_material_rows.side_effect = lambda grid,tasks: ({t.material:[0,1] if duplicate else [0] for t in tasks},{})
        self.grid.ModifyCheckBox.side_effect = lambda *args: self.assertTrue(any(e['type']=='INTERACTION_RESOLVED' for e in self.events))
        waits = []
        def on_event(kind, data):
            nonlocal duplicate
            if kind == 'ACTION_REQUIRED' and data.get('state') == 'WAITING_FOR_USER':
                waits.append(data)
                self.assertEqual(self.grid.ModifyCheckBox.call_count,0)
                if len(waits) == 2:
                    duplicate = False  # Simulate the user's SAP correction.
                self.reply(data)
        def process(sap, excel, group):
            for t in group.tasks: t.buyer_receipt_row = 0
            _, valid = engine.prepare_prod_direct_rfq_rows(sap,self.grid,group.tasks)
            excel.write_status([t.excel_row for t in valid], rfq_status='SUCCESS',rfq_number='MOCK-2')
            excel.save()
            return 'SUCCESS','MOCK-2'
        with patch.object(engine,'SAP_LONG_WAIT_SEC',0):
            self.assertEqual(self.invoke(process,on_event),0)
        self.assertEqual(len(waits),2)
        self.assertNotEqual(waits[0]['requestId'],waits[1]['requestId'])
        self.assertIn('The issue is still present.',waits[1]['message'])
        self.assertEqual(self.sap.press.call_count,1)
        self.sap.press.assert_called_once_with(engine.ID_CREATE_RFQ_FROM_BR)
        self.grid.ModifyCheckBox.assert_called_once_with(0,'LIFNR1_CB',False)
        self.assertEqual(self.events[-1]['type'],'RUN_COMPLETED')

    def test_buyer_manual_saved_green_resume_does_not_save_twice(self):
        self.book()
        self.title = 'Buyer Receipt'
        self.green = False
        self.sap.map_material_rows.return_value = ({'111':[0]}, {})
        waits = []
        original_save = MagicMock(side_effect=engine.SapRpaError('BUYER_RECEIPT_SAVE_POPUP','Technology required: missing master data'))
        def on_event(kind,data):
            if kind=='ACTION_REQUIRED' and data.get('state')=='WAITING_FOR_USER':
                waits.append(data)
                if len(waits)==2: self.green = True
                self.reply(data)
        completed = []
        def process(sap,excel,group):
            accepted,_,_ = engine.save_buyer_receipt(sap,self.grid,group.tasks)
            completed.append([t.material for t in accepted])
            excel.write_status([t.excel_row for t in accepted],buyer_status='SUCCESS',rfq_status='SUCCESS',rfq_number='MOCK-3')
            excel.save()
            return 'SUCCESS','MOCK-3'
        self.assertEqual(self.invoke(process,on_event,original_save),0)
        self.assertEqual(len(waits),2)
        self.assertEqual(original_save.call_count,1)
        self.assertEqual(completed,[['111']])
        self.sap.press.assert_not_called()

    def test_stop_unknown_preserves_created_rfq_and_prevents_next_group(self):
        self.book([['C100','P1','111','26517','a@example.com','2030-12-31'],
                   ['C100','P2','222','26517','a@example.com','2030-12-31'],
                   ['C100','P3','333','26517','a@example.com','2030-12-31']])
        processed = []
        def process(sap,excel,group):
            processed.append(group.project)
            if group.project=='P1':
                excel.write_status([2],rfq_status='SUCCESS',rfq_number='KEEP-1')
                excel.save()
                return 'SUCCESS','KEEP-1'
            excel.write_status([3],buyer_status='SUCCESS',rfq_status='CREATED_NUMBER_NOT_FOUND')
            excel.save()
            raise engine.SapRpaError('RFQ_NUMBER','RFQ outcome uncertain')
        def on_event(kind,data):
            if kind=='ACTION_REQUIRED' and data.get('state')=='WAITING_FOR_USER':
                self.assertEqual(data['allowedActions'],['stop'])
                self.reply(data,'stop')
        self.assertEqual(self.invoke(process,on_event),3)
        self.assertEqual(processed,['P1','P2'])
        statuses = self.statuses()
        self.assertEqual(statuses[0]['RFQ Number'],'KEEP-1')
        self.assertEqual(statuses[0]['RFQ Status'],'SUCCESS')
        self.assertEqual(statuses[1]['RFQ Status'],'CREATED_NUMBER_NOT_FOUND')
        self.assertEqual(statuses[2]['RFQ Status'],'CANCELLED')
        self.assertEqual(self.events[-1]['type'],'RUN_CANCELLED')
        self.assertEqual(self.events[-1]['rfqNumbers'],['KEEP-1'])
        log = next(Path(self.fixture.temp.name).glob('*.csv')).read_text(encoding='utf-8-sig')
        self.assertIn('CANCELLED',log)

    def test_known_create_prerequisite_resumes_only_from_manual_saved_buyer(self):
        self.book()
        self.title = 'Buyer Receipt'
        self.sap.map_material_rows.return_value = ({'111':[0]}, {})
        calls = []
        def process(sap,excel,group):
            calls.append('create-buyer')
            raise engine.SapRpaError('CREATE_BUYER_RECEIPT','Technology required: missing master data')
        def complete(sap,excel,group,grid,tasks):
            calls.append('complete-rfq')
            excel.write_status([t.excel_row for t in tasks],rfq_status='SUCCESS',rfq_number='MOCK-4')
            excel.save()
            return 'SUCCESS','MOCK-4'
        def on_event(kind,data):
            if kind=='ACTION_REQUIRED' and data.get('state')=='WAITING_FOR_USER':
                self.assertEqual(data['allowedActions'],['continue','stop'])
                self.reply(data)
        with patch.object(engine,'complete_rfq_from_buyer_receipt',side_effect=complete):
            self.assertEqual(self.invoke(process,on_event),0)
        self.assertEqual(calls,['create-buyer','complete-rfq'])
        self.sap.press.assert_not_called()

    def test_unknown_buyer_popup_never_offers_continue(self):
        self.book()
        original_save = MagicMock(side_effect=engine.SapRpaError('BUYER_RECEIPT_SAVE_POPUP','Unexpected authorization/process condition'))
        def process(sap,excel,group):
            engine.save_buyer_receipt(sap,self.grid,group.tasks)
        def on_event(kind,data):
            if kind=='ACTION_REQUIRED' and data.get('state')=='WAITING_FOR_USER':
                self.assertEqual(data['allowedActions'],['stop'])
                self.reply(data,'stop')
        self.assertEqual(self.invoke(process,on_event,original_save),3)
        self.assertEqual(original_save.call_count,1)

    def test_stop_waiting_known_checkpoint_does_not_resume(self):
        self.book()
        def on_event(kind,data):
            if kind=='ACTION_REQUIRED' and data.get('state')=='WAITING_FOR_USER':
                self.reply(data,'stop')
        def process(sap,excel,group):
            for task in group.tasks: task.buyer_receipt_row = 0
            engine.prepare_prod_direct_rfq_rows(sap,self.grid,group.tasks)
            self.fail('Stop must not return to execution')
        with patch.object(engine,'SAP_LONG_WAIT_SEC',0):
            self.assertEqual(self.invoke(process,on_event),3)
        self.grid.ModifyCheckBox.assert_not_called()
        self.assertEqual(self.statuses()[0]['RFQ Status'],'CANCELLED')

    def test_revalidation_rejects_changed_session_busy_popup_or_error(self):
        identity = session_identity(self.sap)
        for change in ['session','busy','popup','error']:
            with self.subTest(change=change):
                self.sap.session.Info.Client = '200' if change=='session' else '100'
                self.sap.session.Busy = change=='busy'
                self.sap.exists.return_value = change=='popup'
                self.sap.popup_message_text.return_value = 'Resolve popup'
                self.sap.status_bar.return_value = ('E','Error') if change=='error' else ('S','')
                with self.assertRaises(ValueError): check_session(self.sap,identity)
        self.sap.press.assert_not_called()

    def test_wrong_screen_supplier_or_quantity_cannot_resume(self):
        self.book()
        self.fixture.preview()
        store = engine.ExcelStore(self.fixture.path,'RPA_Input')
        try: tasks = store.load_tasks()
        finally: store.close()
        self.sap.map_material_rows.return_value = ({'111':[0]}, {})
        self.title = 'NPL'
        with self.assertRaises(ValueError): verify_staging(engine,self.sap,tasks)
        self.title = 'RFQ/Quotation Creation by vendor'
        with patch.object(self.sap,'get_cell',return_value='WRONG'):
            with self.assertRaises(ValueError): verify_staging(engine,self.sap,tasks)
        self.title = 'Buyer Receipt'
        original = self.sap.get_cell.side_effect
        self.sap.get_cell.side_effect = lambda grid,row,col: '999' if col=='Z12MRQTY' else original(grid,row,col)
        with self.assertRaises(ValueError): verify_saved_buyer(engine,self.sap,tasks)
        self.grid.ModifyCheckBox.assert_not_called()
        self.sap.press.assert_not_called()


if __name__ == '__main__':
    # Importing the fixture class must not re-run its already-covered test suite.
    unittest.main(defaultTest='RecoveryTests', argv=[sys.argv[0]], verbosity=2)
