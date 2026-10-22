from __future__ import annotations

import json
import os
import shutil
import sys
import subprocess
import time
ENGINE_IMPORT_STARTED = time.perf_counter()
import traceback
import winreg
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path
from typing import Any

import win32com.client

sys.path.insert(0, str(Path(__file__).resolve().parent))
from hub_user_interaction import HubStopped, HubUserInteraction, session_identity, verify_session
from me01_source_table import SourceListTable, parse_date
ENGINE_IMPORT_SECONDS = time.perf_counter() - ENGINE_IMPORT_STARTED

HUB_INTERACTION = None


PLANT = "C100"
SHEET_NAME = "ME01_Input"
EVENT_PREFIX = "ME01_EVENT "
MATERIAL_HEADERS = {"material no.", "material no", "material", "part number"}
PARMA_HEADERS = {"parma", "parma no.", "parma no", "supplier", "supplier code", "intended supplier"}
DUAL_HEADERS = {"dual sourcing", "dual sourcing (yes/no)", "是否双供", "双供"}
RESULT_HEADERS = {"status": "ME01 Status", "error": "Error Message", "updated": "Updated At"}


def emit(event: str, **payload: Any) -> None:
    print(EVENT_PREFIX + json.dumps({"event": event, **payload}, ensure_ascii=False), flush=True)


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_identifier(value: Any) -> str:
    text = clean(value)
    if text.isdigit():
        return text.lstrip("0") or "0"
    return text.upper()


def parse_dual_sourcing(value: Any, row: int) -> bool:
    text = clean(value).lower()
    if text in {"yes", "true", "1", "是", "√", "✓", "☑"}:
        return True
    if text in {"no", "false", "0", "否"}:
        return False
    raise RuntimeError(f"Excel row {row}: Dual Sourcing / 是否双供 is required. Select Yes or No; blank is not No.")


def same_path(first: Any, second: Path) -> bool:
    try:
        return os.path.normcase(os.path.abspath(clean(first))) == os.path.normcase(str(second))
    except Exception:
        return False


def find_open_excel_workbook(path: Path) -> Any | None:
    """Attach only this file, including another Excel instance; never open/close it."""
    readonly = None
    try:
        excel = win32com.client.GetActiveObject("Excel.Application")
        for index in range(1, excel.Workbooks.Count + 1):
            workbook = excel.Workbooks(index)
            if same_path(workbook.FullName, path):
                if not getattr(workbook, 'ReadOnly', False):
                    return workbook
                readonly = workbook
    except Exception:
        pass
    try:
        import pythoncom
        running = pythoncom.GetRunningObjectTable()
        context = pythoncom.CreateBindCtx(0)
        for moniker in running.EnumRunning():
            try:
                name = moniker.GetDisplayName(context, None).lstrip('!')
                if not same_path(name, path):
                    continue
                workbook = win32com.client.Dispatch(running.GetObject(moniker))
                if same_path(workbook.FullName, path):
                    if not getattr(workbook, 'ReadOnly', False):
                        return workbook
                    readonly = workbook
            except Exception:
                continue
    except Exception:
        pass
    return readonly


def excel_call(action):
    """Retry transient Excel-busy errors only, never SAP actions or arbitrary errors."""
    for attempt in range(3):
        try:
            return action()
        except Exception as error:
            if getattr(error, 'hresult', None) not in {-2147418111, -2147417846} or attempt == 2:
                raise
            time.sleep(0.15)


def excel_column(index: int) -> str:
    text = ''
    while index:
        index, remainder = divmod(index - 1, 26)
        text = chr(65 + remainder) + text
    return text


def com_range_values(sheet: Any, first_row: int, last_row: int, columns: int) -> tuple:
    """Read one Excel rectangle, not one COM round-trip per input cell."""
    if first_row < 1 or last_row < first_row or columns < 1:
        raise RuntimeError('Invalid Excel input range.')
    try:
        reader = sheet.Range
    except AttributeError:
        # Compatibility for old adapters; real Excel supports Range.Value2.
        return tuple(tuple(sheet.Cells(row, column).Value for column in range(1, columns + 1))
                     for row in range(first_row, last_row + 1))
    values = reader(f'A{first_row}:{excel_column(columns)}{last_row}').Value2
    if first_row == last_row and columns == 1:
        values = ((values,),)
    if (not isinstance(values, (tuple, list)) or len(values) != last_row-first_row+1
            or any(not isinstance(row, (tuple, list)) or len(row) != columns for row in values)):
        raise RuntimeError('Excel bulk input range is incomplete. No SAP record will be started.')
    return tuple(tuple(row) for row in values)


class WorkbookResults:
    def __init__(self, path: Path, backup_path: Path) -> None:
        self.path = path
        self.backup_path = backup_path
        self.result_path = path
        self.writeback_warning = ''
        self.pending: dict[int, dict[str, str]] = {}
        self.dirty: set[int] = set()
        self.input_rows: dict[int, tuple] = {}
        self.open_excel_workbook = find_open_excel_workbook(path)
        self.workbook: Any
        self.sheet: Any
        self.result_columns: dict[str, int] = {}
        self.assignments: OrderedDict[str, dict[str, Any]] = OrderedDict()
        if self.open_excel_workbook is not None:
            self.mode = "excel-com"
            self.workbook = self.open_excel_workbook
            excel_call(lambda: self.workbook.SaveCopyAs(str(backup_path)))
            self.sheet = self._excel_sheet()
            try:
                excel_call(self._read_excel_com)
            except RuntimeError:
                raise  # Invalid template/input is not an Excel-lock problem.
            except Exception as error:
                # Even a protected/busy workbook can be kept open. Validate the
                # exact SaveCopyAs snapshot and keep all outputs in owned logs.
                self._save_result_copy(error)
                self.assignments, self.input_rows = OrderedDict(), {}
                self._read_openpyxl()
            if self.mode == 'excel-com':
                emit("preparing", message="The workbook is open in Excel. Results will be written through Excel without closing it.")
        else:
            # Do not pay openpyxl's import/load cost for an already-open workbook.
            import openpyxl
            self.mode = "openpyxl"
            shutil.copy2(path, backup_path)
            self.workbook = openpyxl.load_workbook(path, keep_vba=path.suffix.lower() == '.xlsm')
            self.sheet = self._input_sheet(list(self.workbook.worksheets), lambda sheet: [clean(cell.value).lower() for cell in sheet[1]])
            self._read_openpyxl()
        if not self.assignments:
            raise RuntimeError("No Material No. and Parma values were found in the uploaded workbook.")

    def _input_sheet(self, sheets, header_values):
        candidates = [sheet for sheet in sheets if MATERIAL_HEADERS.intersection(header_values(sheet)) and PARMA_HEADERS.intersection(header_values(sheet))]
        for sheet in candidates:
            if clean(getattr(sheet, "Name", getattr(sheet, "title", ""))).lower() == SHEET_NAME.lower():
                return sheet
        if len(candidates) != 1:
            raise RuntimeError("Keep one Source List input worksheet containing Material No. and Parma.")
        return candidates[0]

    def _excel_sheet(self) -> Any:
        try:
            preferred = self.workbook.Worksheets(SHEET_NAME)
            titles = [clean(value).lower() for value in com_range_values(preferred, 1, 1, max(1, int(preferred.UsedRange.Columns.Count)))[0]]
            if MATERIAL_HEADERS.intersection(titles) and PARMA_HEADERS.intersection(titles):
                return preferred
        except Exception:
            pass
        sheets = [self.workbook.Worksheets(index) for index in range(1, self.workbook.Worksheets.Count + 1)]
        return self._input_sheet(sheets, lambda sheet: [clean(value).lower() for value in com_range_values(sheet, 1, 1, max(1, int(sheet.UsedRange.Columns.Count)))[0]])

    def _add_assignment(self, material: str, parma: str, row: int, dual_value: Any) -> None:
        if not material and not parma and not clean(dual_value):
            return
        if not material or not parma:
            raise RuntimeError(f"Excel row {row} must contain both Material No. and Parma.")
        material_key = normalize_identifier(material)
        parma_key = normalize_identifier(parma)
        dual = parse_dual_sourcing(dual_value, row)
        self.input_rows[row] = (material_key, parma_key, dual)
        existing = self.assignments.get(material_key)
        if existing and normalize_identifier(existing["parma"]) != parma_key:
            raise RuntimeError(f"Material {material} has more than one Parma. Keep only the supplier that should be fixed.")
        if existing:
            if existing["dualSourcing"] != dual:
                raise RuntimeError(f"Material {material} has conflicting Dual Sourcing values. Use the same Yes/No for duplicate rows.")
            existing["rows"].append(row)
        else:
            self.assignments[material_key] = {"material": material, "parma": parma, "dualSourcing": dual, "rows": [row]}

    def _read_excel_com(self) -> None:
        # A transient Excel-busy retry restarts parsing, never duplicates rows.
        self.assignments, self.input_rows, self.result_columns = OrderedDict(), {}, {}
        max_column = max(1, int(self.sheet.UsedRange.Columns.Count))
        max_row = max(1, int(self.sheet.UsedRange.Rows.Count))
        material_column = 0
        parma_column = 0
        dual_column = 0
        headers = com_range_values(self.sheet, 1, 1, max_column)[0]
        for column, value in enumerate(headers, start=1):
            title = clean(value).lower()
            if title in MATERIAL_HEADERS:
                material_column = column
            if title in PARMA_HEADERS:
                parma_column = column
            if title in DUAL_HEADERS:
                dual_column = column
            for key, result_title in RESULT_HEADERS.items():
                if title == result_title.lower():
                    self.result_columns[key] = column
        if not material_column or not parma_column or not dual_column:
            raise RuntimeError('Required columns: Material No., Parma, Dual Sourcing / 是否双供. Download the new ME01 template.')
        self.input_columns = (material_column, parma_column, dual_column)
        for key, title in RESULT_HEADERS.items():
            if key not in self.result_columns:
                max_column += 1
                self.sheet.Cells(1, max_column).Value = title
                self.result_columns[key] = max_column
        input_columns = max(material_column, parma_column, dual_column)
        block_rows = min(4096, max(1, 100_000 // input_columns))
        for first in range(2, max_row + 1, block_rows):
            block = com_range_values(self.sheet, first, min(max_row, first + block_rows - 1), input_columns)
            for offset, values in enumerate(block):
                self._add_assignment(clean(values[material_column-1]), clean(values[parma_column-1]),
                                     first+offset, values[dual_column-1])

    def _read_openpyxl(self) -> None:
        material_column = 0
        parma_column = 0
        dual_column = 0
        for column in range(1, self.sheet.max_column + 1):
            title = clean(self.sheet.cell(1, column).value).lower()
            if title in MATERIAL_HEADERS:
                material_column = column
            if title in PARMA_HEADERS:
                parma_column = column
            if title in DUAL_HEADERS:
                dual_column = column
            for key, result_title in RESULT_HEADERS.items():
                if title == result_title.lower():
                    self.result_columns[key] = column
        if not material_column or not parma_column or not dual_column:
            raise RuntimeError('Required columns: Material No., Parma, Dual Sourcing / 是否双供. Download the new ME01 template.')
        for key, title in RESULT_HEADERS.items():
            if key not in self.result_columns:
                column = self.sheet.max_column + 1
                self.sheet.cell(1, column, title)
                self.result_columns[key] = column
        for row in range(2, self.sheet.max_row + 1):
            self._add_assignment(clean(self.sheet.cell(row, material_column).value), clean(self.sheet.cell(row, parma_column).value), row, self.sheet.cell(row, dual_column).value)

    def write(self, rows: list[int], status: str, message: str) -> None:
        timestamp = datetime.now().astimezone().isoformat(timespec="seconds")
        for row in rows:
            values = {"status": status, "error": message, "updated": timestamp}
            self.pending[row] = values
            self.dirty.add(row)

    def _flush_excel(self) -> None:
        if getattr(self.workbook, 'ReadOnly', False):
            raise PermissionError('Excel workbook is read-only; original results cannot be saved.')
        for column, allowed in zip(self.input_columns, (MATERIAL_HEADERS, PARMA_HEADERS, DUAL_HEADERS)):
            if clean(self.sheet.Cells(1, column).Value).lower() not in allowed:
                raise RuntimeError('Excel input columns moved during the run; original columns will not be overwritten.')
        for key, column in self.result_columns.items():
            if clean(self.sheet.Cells(1, column).Value).lower() != RESULT_HEADERS[key].lower():
                raise RuntimeError('Excel result columns moved during the run; original columns will not be overwritten.')
        for row in self.dirty:
            material, parma, dual = (self.sheet.Cells(row, column).Value for column in self.input_columns)
            if (normalize_identifier(material), normalize_identifier(parma), parse_dual_sourcing(dual, row)) != self.input_rows[row]:
                raise RuntimeError(f'Excel row {row} changed while SAP was running. Original input will not be overwritten.')
        for row in self.dirty:
            for key, value in self.pending[row].items():
                self.sheet.Cells(row, self.result_columns[key]).Value = value
        self.workbook.Save()

    def _attach_late_excel(self, workbook) -> None:
        previous = self.workbook
        original = (self.workbook, self.sheet, self.assignments, self.input_rows, self.result_columns, getattr(self, 'input_columns', None))
        try:
            self.workbook = workbook
            self.sheet = self._excel_sheet()
            excel_call(self._read_excel_com)
            if self.input_rows != original[3]:
                raise RuntimeError('Excel inputs changed after the batch started. Results will not overwrite the changed workbook.')
            self.assignments, self.input_rows = original[2:4]
        except Exception:
            (self.workbook, self.sheet, self.assignments, self.input_rows, self.result_columns, self.input_columns) = original
            raise
        self.mode = 'excel-com'
        self.open_excel_workbook = workbook
        self.dirty.update(self.pending)
        previous.close()

    def _save_result_copy(self, reason) -> None:
        import openpyxl
        # Preserve completed SAP results independently of Excel locks/COM. Never
        # reopen or replay a SAP record because a local workbook could not save.
        previous, previous_mode = self.workbook, self.mode
        workbook = openpyxl.load_workbook(self.backup_path, keep_vba=self.path.suffix.lower() == '.xlsm')
        sheet = self._input_sheet(list(workbook.worksheets), lambda item: [clean(cell.value).lower() for cell in item[1]])
        headers = {clean(cell.value).lower(): index for index, cell in enumerate(sheet[1], 1)}
        self.result_columns = {}
        for key, title in RESULT_HEADERS.items():
            column = headers.get(title.lower()) or sheet.max_column + 1
            sheet.cell(1, column, title)
            self.result_columns[key] = column
        for row, values in self.pending.items():
            for key, value in values.items():
                sheet.cell(row, self.result_columns[key], value)
        target = self.backup_path.parent / f'{self.path.stem}_ME01_results{self.path.suffix}'
        workbook.save(target)
        self.workbook, self.sheet, self.mode, self.result_path = workbook, sheet, 'result-copy', target
        if previous_mode == 'openpyxl':
            previous.close()
        self.writeback_warning = (f'Original Excel could not be saved: {clean(reason)}. Excel may stay open. '
                                  f'Completed SAP results are preserved here; do not rerun SAP: {target}')
        emit('warning', message=self.writeback_warning, resultPath=str(target))

    def save(self) -> None:
        if not self.dirty:
            return
        try:
            if self.mode == 'openpyxl':
                live = find_open_excel_workbook(self.path)
                if live is not None:
                    self._attach_late_excel(live)
            if self.mode == 'excel-com':
                excel_call(self._flush_excel)
            else:
                for row in self.dirty:
                    for key, value in self.pending[row].items():
                        self.sheet.cell(row, self.result_columns[key], value)
                self.workbook.save(self.result_path)
        except Exception as error:
            if self.mode == 'result-copy':
                raise  # Even the owned log folder is unwritable: stop, never replay SAP.
            self._save_result_copy(error)
        self.dirty.clear()


def wait_ready(session: Any, timeout: float = 30.0) -> None:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        try:
            if not session.Busy:
                return
        except Exception:
            pass
        time.sleep(0.15)
    raise RuntimeError("SAP GUI did not become ready within 30 seconds.")


def status_error(session: Any) -> str:
    try:
        bar = session.findById("wnd[0]/sbar")
        if str(bar.MessageType).upper() in {"E", "A"}:
            return clean(bar.Text) or "SAP returned an error message."
    except Exception:
        return ""
    return ""


def try_active_session() -> Any | None:
    try:
        application = win32com.client.GetObject("SAPGUI").GetScriptingEngine
    except Exception:
        return None
    try:
        for connection_index in range(application.Children.Count):
            connection = application.Children(connection_index)
            for session_index in range(connection.Children.Count):
                session = connection.Children(session_index)
                session.findById("wnd[0]")
                return session
    except Exception:
        return None
    return None


def sap_logon_candidates() -> list[Path]:
    candidates: list[Path] = []
    configured = os.environ.get("SAPLOGON_EXE", "").strip()
    if configured:
        candidates.append(Path(configured))
    for hive in (winreg.HKEY_CURRENT_USER, winreg.HKEY_LOCAL_MACHINE):
        for registry_path in (
            r"SOFTWARE\Microsoft\Windows\CurrentVersion\App Paths\saplogon.exe",
            r"SOFTWARE\WOW6432Node\Microsoft\Windows\CurrentVersion\App Paths\saplogon.exe",
        ):
            try:
                with winreg.OpenKey(hive, registry_path) as key:
                    candidates.append(Path(clean(winreg.QueryValue(key, None))))
            except OSError:
                pass
    for root in (os.environ.get("ProgramFiles(x86)"), os.environ.get("ProgramFiles")):
        if root:
            candidates.append(Path(root) / "SAP" / "FrontEnd" / "SAPgui" / "saplogon.exe")
    return candidates


def launch_sap_logon() -> None:
    executable = next((path for path in sap_logon_candidates() if path.is_file()), None)
    if executable is None:
        raise RuntimeError("No SAP GUI session is open and saplogon.exe could not be found. Open SAP Logon manually, then retry.")
    subprocess.Popen([str(executable)], close_fds=True)


def find_active_session() -> Any:
    session = try_active_session()
    if session is not None:
        return session
    launch_sap_logon()
    emit("waiting-for-sap", message="SAP Logon was opened. Select the correct system and complete sign-in; waiting up to 3 minutes.")
    deadline = time.monotonic() + 180
    while time.monotonic() < deadline:
        session = try_active_session()
        if session is not None:
            return session
        time.sleep(1)
    raise RuntimeError("SAP Logon was opened, but no signed-in SAP GUI session became available within 3 minutes.")


def find_hub_session(connector=None) -> Any:
    from sap_gui_connection import SapGuiConnector
    if connector is None:
        connector = SapGuiConnector("VCE - One Digital Core [949]", "VCE", "100",
            launch=launch_sap_logon, stop=lambda: bool(HUB_INTERACTION and HUB_INTERACTION.stop_event.is_set()),
            cancelled=HubStopped)
    # A verified existing business window may be reused. open_me01 navigates
    # once; SAP save/discard and permission prompts remain manual.
    return connector.select()[0]


def open_me01(session: Any) -> None:
    session.findById("wnd[0]").maximize()
    command = session.findById("wnd[0]/tbar[0]/okcd")
    command.Text = "/nME01"
    session.findById("wnd[0]").sendVKey(0)
    wait_ready(session)
    error = status_error(session)
    if error:
        raise RuntimeError(f"ME01 could not be opened: {error}")


def source_date(value: Any) -> date:
    return parse_date(value)


def check_source_stop() -> None:
    if HUB_INTERACTION and HUB_INTERACTION.stop_event.is_set():
        raise HubStopped()


def write_run_diagnostic(directory: Path, diagnostics: dict) -> None:
    # Audit failures must never cause SAP actions/Save to be replayed.
    try:
        with open(directory / 'me01-performance.jsonl', 'a', encoding='utf-8') as log:
            log.write(json.dumps(diagnostics, ensure_ascii=False) + '\n')
    except OSError:
        pass


def maintain_material(session: Any, material: str, parma: str, dual_sourcing: bool,
                      valid_from: str, valid_to: str, purchasing_org: str, diagnostics: dict | None = None) -> bool:
    started = time.perf_counter()
    material_field = session.findById("wnd[0]/usr/ctxtEORD-MATNR")
    plant_field = session.findById("wnd[0]/usr/ctxtEORD-WERKS")
    material_field.Text = material
    plant_field.Text = PLANT
    plant_field.SetFocus()
    plant_field.caretPosition = len(PLANT)
    session.findById("wnd[0]").sendVKey(0)
    wait_ready(session)
    error = status_error(session)
    if error:
        raise RuntimeError(error)
    identity = session_identity(session) if HUB_INTERACTION else None
    def verify_context():
        if identity:
            verify_session(session, identity)
        if (normalize_identifier(session.findById('wnd[0]/usr/ctxtEORD-MATNR').Text) != normalize_identifier(material)
                or clean(session.findById('wnd[0]/usr/ctxtEORD-WERKS').Text).upper() != PLANT):
            raise RuntimeError('The SAP material/Plant changed. No row will be guessed; return to the original record.')
    adapter = SourceListTable(session, "wnd[0]/usr/tblSAPLMEORTC_0205", wait_ready, check_source_stop, verify_context)
    navigation_finished = time.perf_counter()
    save_started = None
    saved = False
    try:
        changed = adapter.upsert(parma, dual_sourcing, valid_from, valid_to, purchasing_org, skip_existing=True)
        if not changed:
            if diagnostics is not None:
                diagnostics['skipReason'] = adapter.skip_reason
            # This branch has performed no data edits or Save. Start a fresh
            # ME01 selection so the next material cannot reuse this overview.
            check_source_stop()
            if session.findById('wnd[1]', False) is not None:
                raise RuntimeError('An SAP dialog is open. The existing record was not edited; handle the dialog manually.')
            open_me01(session)
            return False
        adapter.stage = "Save Source List"
        save_started = time.perf_counter()
        session.findById("wnd[0]/tbar[0]/btn[11]").Press()
        wait_ready(session)
        error = status_error(session)
        if error:
            raise RuntimeError(error)
        if session.findById("wnd[1]", False) is not None or clean(session.findById("wnd[0]/sbar").MessageType).upper() != "S":
            raise RuntimeError("SAP Save has no confirmed success response or has an open dialog. Stop and review SAP; Save will not be replayed.")
        saved = True
        return True
    except Exception as error:
        raise RuntimeError(f"[{adapter.stage}] {clean(error)}") from error
    finally:
        if diagnostics is not None:
            finished = time.perf_counter()
            diagnostics.update({
                'sapNavigationSeconds': round(navigation_finished-started, 3),
                'tableMaintenanceSeconds': round((save_started or finished)-navigation_finished, 3),
                'sapSaveSeconds': round(finished-save_started, 3) if save_started is not None else 0,
                'sapTotalSeconds': round(finished-started, 3),
                'savedConfirmed': saved, 'lastStage': adapter.stage,
                'reader': dict(adapter.metrics), 'readTrace': adapter.read_trace,
                'existingRowsBefore': adapter.before_rows, 'plannedRowsAfter': adapter.expected_rows,
                'existingRowPolicy': 'target supplier exists: skip all changes; otherwise retain rows and dates, only other Fix / MRP may change'
            })


def main() -> int:
    global HUB_INTERACTION
    startup_started = time.perf_counter()
    HUB_INTERACTION = HubUserInteraction("Source List", emit) if os.environ.get("HUB_RUN_ID") else None
    if HUB_INTERACTION:
        HUB_INTERACTION.start_listener()
    valid_from = os.environ.get('ME01_VALID_FROM', date.today().isoformat())
    valid_to = os.environ.get('ME01_VALID_TO', '2099-12-31')
    purchasing_org = os.environ.get('ME01_PURCH_ORG', 'C100').strip().upper()
    if parse_date(valid_from) > parse_date(valid_to) or not purchasing_org or len(purchasing_org) > 4 or not purchasing_org.isalnum():
        raise RuntimeError('Invalid Source List validity or purchasing organization.')
    excel_path = Path(os.environ.get("EXCEL_PATH", "")).expanduser().resolve()
    if not excel_path.is_file():
        raise RuntimeError("The selected ME01 workbook was not found.")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    from hub_artifacts import artifact_directory
    run_artifacts = artifact_directory('me01')
    backup_path = run_artifacts / f"{excel_path.stem}_backup_{timestamp}{excel_path.suffix}"
    workbook_started = time.perf_counter()
    results = WorkbookResults(excel_path, backup_path)
    workbook_seconds = time.perf_counter() - workbook_started
    emit("preparing", message=f"Backup created: {backup_path}", backupPath=str(backup_path))

    total = len(results.assignments)
    emit("connecting", message="Connecting to SAP GUI. SAP Logon will open automatically if needed.", total=total)
    succeeded = failed = skipped = 0
    cancelled = False
    try:
        connect_started = time.perf_counter()
        if HUB_INTERACTION:
            from sap_gui_connection import SapGuiConnector
            connector = SapGuiConnector("VCE - One Digital Core [949]", "VCE", "100",
                launch=launch_sap_logon, stop=HUB_INTERACTION.stop_event.is_set, cancelled=HubStopped)
            verify_startup = lambda: find_hub_session(connector)
            try:
                session = verify_startup()
            except Exception as error:
                session = HUB_INTERACTION.wait("SIGN_IN",
                    getattr(error, 'issue_summary', "SAP startup verification failed. Review the details, then Check Again."),
                    getattr(error, 'issue_summary_zh', "SAP 启动检查未通过，请查看具体原因后重新检查。"),
                    verify_startup, detail=str(error))
        else:
            session = find_active_session()
        identity = session_identity(session) if HUB_INTERACTION else None
        connect_seconds = time.perf_counter() - connect_started
        navigation_started = time.perf_counter()
        open_me01(session)
        write_run_diagnostic(run_artifacts, {
            'phase': 'startup', 'timestamp': datetime.now().astimezone().isoformat(),
            'importSeconds': round(ENGINE_IMPORT_SECONDS, 3),
            'workbookSeconds': round(workbook_seconds, 3), 'workbookMode': results.mode,
            'sapConnectSeconds': round(connect_seconds, 3),
            'openMe01Seconds': round(time.perf_counter()-navigation_started, 3),
            'engineReadySeconds': round(time.perf_counter()-startup_started+ENGINE_IMPORT_SECONDS, 3),
            'recordCount': total,
        })
        for current, assignment in enumerate(results.assignments.values(), start=1):
            if HUB_INTERACTION and HUB_INTERACTION.stop_event.is_set():
                raise HubStopped()
            material = assignment["material"]
            parma = assignment["parma"]
            excel_rows = assignment["rows"]
            dual = assignment["dualSourcing"]
            record_started = time.perf_counter()
            diagnostics = {'material': material, 'parma': parma, 'excelRows': excel_rows, 'dualSourcing': dual}
            outcome = 'failed'
            emit("record", message=f"Material {material}: locating Parma {parma} in Plant {PLANT}.",
                 material=material, parma=parma, current=current, total=total, status="running")
            try:
                if HUB_INTERACTION:
                    verify_session(session, identity)
                changed = maintain_material(session, material, parma, dual, valid_from, valid_to, purchasing_org, diagnostics)
                outcome = 'saved' if changed else 'already-exists'
                if changed:
                    succeeded += 1
                else:
                    skipped += 1
                results.write(excel_rows, "SUCCESS" if changed else "SKIPPED", "" if changed else "SOURCE_EXISTS: target supplier already exists; no SAP changes.")
                mode = "Yes: other Fix / MRP retained" if dual else "No: other Fix / MRP cleared"
                elapsed = time.perf_counter()-record_started
                explanation = f'Dual Sourcing {mode}: verified and saved' if changed else 'supplier already exists; skipped without changes; continuing to next material'
                emit("record", message=f"{material} / Parma {parma}: {explanation}. {elapsed:.1f}s.",
                     material=material, parma=parma, current=current, total=total, status="success" if changed else "skipped")
            except HubStopped:
                outcome = 'cancelled'
                raise
            except Exception as error:
                failed += 1
                message = clean(error) or error.__class__.__name__
                try:
                    with open(run_artifacts / 'me01-errors.log', 'a', encoding='utf-8') as log:
                        log.write(f'Excel rows {excel_rows}; Material {material}; Parma {parma}; Dual Sourcing {dual}\n{traceback.format_exc()}\n')
                except OSError as log_error:
                    message += f'; technical log could not be written: {clean(log_error)}'
                results.write(excel_rows, "ERROR", message)
                emit("record", message=f"{material} / Parma {parma}: {message}", material=material,
                     parma=parma, current=current, total=total, status="failed")
                if HUB_INTERACTION:
                    results.save()
                    HUB_INTERACTION.wait("UNVERIFIED_STATE",
                        f"Material {material} / Parma {parma} could not be completed. Stop and review SAP; the detailed error is below. Save will not be replayed.",
                        f"物料 {material} / Parma {parma} 未完成：{message}。请检查 SAP，不会自动重复保存。",
                        detail=message,
                        materials=[material], rows=excel_rows,
                        instructions=f'Technical trace: {run_artifacts / "me01-errors.log"}',
                        instructionsZh=f'完整错误记录：{run_artifacts / "me01-errors.log"}')
                else:
                    raise  # Unknown/partially edited state must not be discarded by /nME01.
            finally:
                diagnostics['sapAndResultWriteSeconds'] = round(time.perf_counter()-record_started, 3)
                diagnostics['outcome'] = outcome
                diagnostics['timestamp'] = datetime.now().astimezone().isoformat()
                # Failed/stopped records cannot reach the normal completion block.
                if outcome in {'failed', 'cancelled'}:
                    write_run_diagnostic(run_artifacts, diagnostics)
            excel_save_started = time.perf_counter()
            results.save()
            diagnostics['excelSaveSeconds'] = round(time.perf_counter()-excel_save_started, 3)
            diagnostics['totalRecordSeconds'] = round(time.perf_counter()-record_started, 3)
            diagnostics['targetSeconds'] = 20
            diagnostics['overTarget'] = diagnostics['totalRecordSeconds'] > 20
            if outcome in {'saved', 'already-exists'}:
                write_run_diagnostic(run_artifacts, diagnostics)
            emit('performance', material=material, elapsedSeconds=diagnostics['totalRecordSeconds'],
                 excelSaveSeconds=diagnostics['excelSaveSeconds'], targetSeconds=20, overTarget=diagnostics['overTarget'])

    except HubStopped:
        cancelled = True
    finally:
        results.save()
    emit("complete", message="Source List stopped safely." if cancelled else "ME01 source-list batch finished.",
         processed=succeeded + failed + skipped, succeeded=succeeded, skipped=skipped, failed=failed,
         cancelled=cancelled, resultPath=str(results.result_path), backupPath=str(backup_path), warning=results.writeback_warning)
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as error:
        emit("fatal", message=clean(error) or error.__class__.__name__)
        raise SystemExit(1)
