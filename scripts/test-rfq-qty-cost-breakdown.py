"""Offline regression: quantities, No clears Cost Breakdown; Yes preserves."""
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch, call
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'resources/rpa'))
sys.argv = [sys.argv[0], '--hub', '--validate-only']
import rfq_engine as engine
from rfq_runtime import validate_workbook


class QuantityCostTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name)/'input.xlsx'
        self.book = openpyxl.Workbook()
        self.sheet = self.book.active
        self.sheet.title = 'RPA_Input'
        self.sheet.append(['Plant','Project No.','Material No.','Intended Supplier','Supplier Email','Quotation Due Date',
                           '12 MR Qty','RFQ Qty Prototype','RFQ Qty Serial','Cost Breakdown','PPAP Date','Technology'])

    def tearDown(self):
        self.book.close()
        self.temp.cleanup()

    def tasks(self, choices=('No', 'Yes')):
        for index, choice in enumerate(choices):
            self.sheet.append(['C100','P1',str(100+index),'12345','a@example.com','2030-12-31',120+index,0,25+index,choice,'2030-12-31','TEST_TECH'])
        self.book.save(self.path)
        preview = validate_workbook(engine, self.path)
        store = engine.ExcelStore(self.path, 'RPA_Input')
        tasks = store.load_tasks()
        store.workbook.close()
        for index, task in enumerate(tasks):
            task.buyer_receipt_row = index
        return tasks, preview

    def staging(self, tasks):
        sap, grid = MagicMock(), MagicMock()
        sap.status_bar.return_value = ('S', '')
        sap.exists.return_value = False
        sap.find.return_value = grid
        sap.column_exists.return_value = True
        sap.map_material_rows.return_value = ({t.material:[10-i] for i,t in enumerate(tasks)}, {})
        grid.GetCellType.return_value = 'Checkbox'
        grid.GetCellChangeable.return_value = True
        return sap, grid

    def test_download_template_has_four_fields_and_defaults(self):
        wb = openpyxl.load_workbook(ROOT/'resources/templates/Create_RFQ_Template.xlsx')
        try:
            ws = wb['RPA_Input']
            self.assertEqual([ws.cell(1,c).value for c in range(9,13)],
                             ['12 MR Qty','RFQ Qty Prototype','RFQ Qty Serial','Cost Breakdown'])
            self.assertEqual(ws['D1'].value, 'Supplier Parma')
            self.assertEqual([ws.cell(2,c).value for c in range(9,13)], [None,None,None,'No'])
            for row in range(2,202):
                self.assertEqual([ws.cell(row,c).value for c in range(9,12)], [None,None,None])
            self.assertEqual(ws.tables['CreateRfqInput'].ref,'A1:L201')
            self.assertEqual([c.name for c in ws.tables['CreateRfqInput'].tableColumns],
                             [ws.cell(1,c).value for c in range(1,13)])
            self.assertTrue(any('L2:L201' in str(d.sqref) and 'Yes' in d.formula1 for d in ws.data_validations.dataValidation))
        finally:
            wb.close()

    def test_quantities_and_bool_reach_preview_and_tasks(self):
        tasks, preview = self.tasks(('nO','yEs'))
        self.assertEqual((tasks[0].qty_12mr,tasks[0].rfq_qty_p,tasks[0].rfq_qty_s),('120','0','25'))
        self.assertEqual(preview['sample'][1]['qty12mr'], '121')
        self.assertEqual(preview['sample'][1]['rfqQtySerial'], '26')
        self.assertEqual([r['costBreakdown'] for r in preview['sample']], [False,True])
        self.assertEqual(preview['groupCount'],1)

    def test_required_raw_inputs_cannot_fall_back_to_default(self):
        self.tasks(('No',))
        for col, label in [(7,'12 MR Qty'),(9,'RFQ Qty Serial'),(11,'PPAP Date'),(12,'Technology')]:
            original = self.sheet.cell(2,col).value
            self.sheet.cell(2,col).value = None
            self.book.save(self.path)
            preview = validate_workbook(engine,self.path)
            self.assertEqual(preview['invalidRows'],1,label)
            self.assertIn(label,preview['sample'][0]['message'])
            self.sheet.cell(2,col).value = original

    def test_optional_j_blank_does_not_write_quantity_in_either_grid(self):
        self.tasks(('No',))
        self.sheet.cell(2,8).value = None
        self.book.save(self.path)
        preview = validate_workbook(engine,self.path)
        self.assertEqual(preview['validRows'],1)
        self.assertEqual(preview['sample'][0]['rfqQtyPrototype'],'')
        store = engine.ExcelStore(self.path,'RPA_Input')
        tasks = store.load_tasks()
        store.workbook.close()
        sap,grid = self.staging(tasks)
        engine.fill_buyer_receipt_rows(sap,grid,tasks)
        engine.fill_rfq_grid_rows(sap,grid,tasks)
        self.assertFalse(any(c.args[2] == engine.COL_RFQ_QTY_PROTOTYPE for c in sap.modify_cell.call_args_list))

    def test_missing_new_required_column_is_rejected(self):
        self.tasks(('No',))
        self.sheet.cell(1,12).value = 'Not Technology'
        self.book.save(self.path)
        with self.assertRaises(ValueError):
            validate_workbook(engine,self.path)

    def test_buyer_receipt_quantities_follow_material_mapping(self):
        tasks,_ = self.tasks()
        sap,grid = self.staging(tasks)
        engine.fill_buyer_receipt_rows(sap, grid, tasks)
        for i,task in enumerate(tasks):
            for col,value in [(engine.COL_12MR_QTY,task.qty_12mr),(engine.COL_RFQ_QTY_PROTOTYPE,task.rfq_qty_p),(engine.COL_RFQ_QTY_SERIAL,task.rfq_qty_s)]:
                self.assertIn(call(grid,10-i,col,value),sap.modify_cell.call_args_list)

    def test_prod_no_blank_false_clear_existing_checks(self):
        tasks,_ = self.tasks(('No','','False'))
        sap,grid = self.staging(tasks)
        # No, blank and False clear pre-checked boxes for populated suppliers.
        grid.GetCellCheckBoxChecked.return_value = True
        _, valid = engine.prepare_prod_direct_rfq_rows(sap, object(), tasks)
        self.assertEqual(len(valid),3)
        self.assertEqual(grid.ModifyCheckBox.call_args_list,[call(10,'LIFNR1_CB',False),call(9,'LIFNR1_CB',False),call(8,'LIFNR1_CB',False)])
        self.assertEqual(grid.TriggerModified.call_count,3)
        grid.GetCellType.assert_called()

    def test_prod_mixed_rows_and_suppliers_clear_no_only(self):
        tasks,_ = self.tasks()
        tasks[0].vendors[2] = '33333'
        tasks[0].cost_breakdown[2] = True
        tasks[1].vendors[1] = '22222'
        sap,grid = self.staging(tasks)
        engine.prepare_prod_direct_rfq_rows(sap,object(),tasks)
        self.assertEqual(grid.ModifyCheckBox.call_args_list,[call(10,'LIFNR1_CB',False),call(9,'LIFNR2_CB',False)])

    def test_prod_yes_never_writes_or_triggers(self):
        tasks,_ = self.tasks(('Yes',))
        sap,grid = self.staging(tasks)
        for current in (False,True):
            grid.GetCellCheckBoxChecked.return_value = current
            engine.prepare_prod_direct_rfq_rows(sap,object(),tasks)
        grid.ModifyCheckBox.assert_not_called()
        grid.TriggerModified.assert_not_called()

    def test_wrong_screen_stops_without_checkbox_writes(self):
        tasks,_ = self.tasks(('No',))
        sap,grid = self.staging(tasks)
        grid.GetCellType.return_value = 'Normal'
        with patch.object(engine,'SAP_LONG_WAIT_SEC',0.002), patch.object(engine,'SAP_POLL_SEC',0), self.assertRaises(engine.SapRpaError):
            engine.prepare_prod_direct_rfq_rows(sap,object(),tasks)
        grid.ModifyCheckBox.assert_not_called()

    def test_prod_yes_readonly_checkbox_is_left_unchanged(self):
        tasks,_ = self.tasks(('Yes',))
        sap,grid = self.staging(tasks)
        grid.GetCellChangeable.return_value = False
        _, valid = engine.prepare_prod_direct_rfq_rows(sap,object(),tasks)
        self.assertEqual(valid,tasks)
        grid.GetCellChangeable.assert_not_called()
        grid.ModifyCheckBox.assert_not_called()
        grid.TriggerModified.assert_not_called()

    def test_regular_rfq_grid_clears_no_only_and_preserves_quantities(self):
        tasks,_ = self.tasks()
        sap,grid = self.staging(tasks)
        engine.fill_rfq_grid_rows(sap,grid,tasks)
        cb_calls = [c for c in sap.modify_checkbox.call_args_list if str(c.args[2]).endswith('_CB')]
        self.assertEqual(cb_calls,[call(grid,10,'LIFNR1_CB',False)])
        self.assertIn(call(grid,10,engine.COL_RFQ_QTY_PROTOTYPE,'0'),sap.modify_cell.call_args_list)
        self.assertIn(call(grid,9,engine.COL_RFQ_QTY_SERIAL,'26'),sap.modify_cell.call_args_list)


if __name__ == '__main__':
    unittest.main(argv=[sys.argv[0]],verbosity=2)
