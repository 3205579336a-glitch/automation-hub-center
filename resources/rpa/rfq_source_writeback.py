"""Mirror durable RFQ results to the selected workbook, never replay SAP work.

Excel COM preserves the original workbook's formatting, macros and other sheets.
Only an Excel instance created here may be closed; user-owned workbooks stay open.
"""
from __future__ import annotations

from datetime import date, datetime, time
import os
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

SOURCE_WRITEBACK_POLICY = 'verified-result-columns-v1'


def path_key(path):
    return os.path.normcase(os.path.abspath(str(path)))


def comparable(value, epoch):
    if value is None or value == '':
        return None
    if isinstance(value, (datetime, date, time)):
        return to_excel(value, epoch)
    return value


class ExcelFormula(str):
    """Distinguish real formulas from literal SAP text beginning with '='."""


def row_values(sheet, row, columns):
    def flatten(block):
        if columns == 1 and not isinstance(block, (tuple, list)):
            return [block]
        if not isinstance(block, (tuple, list)) or not block:
            raise ValueError('Excel returned an unreadable row.')
        values = block[0] if isinstance(block[0], (tuple, list)) else block
        if len(values) != columns:
            raise ValueError('Excel returned an incomplete row.')
        return list(values)

    target = sheet.Range(sheet.Cells(row, 1), sheet.Cells(row, columns))
    # Some Office versions stringify every numeric/date constant in Formula.
    # Value2 retains its type; use Formula only for actual formula cells.
    values, formulas = flatten(target.Value2), flatten(target.Formula)
    for index, formula in enumerate(formulas):
        if isinstance(formula, str) and formula.startswith('=') and sheet.Cells(row, index + 1).HasFormula:
            values[index] = ExcelFormula(formula)
    return values


def attach_workbook(source):
    """Find the exact workbook, including a second Excel process, without opening it."""
    import win32com.client

    wanted = path_key(source)
    try:
        app = win32com.client.GetActiveObject('Excel.Application')
        for index in range(1, int(app.Workbooks.Count) + 1):
            book = app.Workbooks.Item(index)
            if path_key(book.FullName) == wanted:
                return app, book
    except Exception:
        pass
    # GetActiveObject exposes only one Excel process. The ROT lets us attach
    # an already-open exact file in another instance; GetObject(file) could
    # instead open a new workbook and is deliberately not used.
    try:
        import pythoncom
        table = pythoncom.GetRunningObjectTable()
        context = pythoncom.CreateBindCtx(0)
        for moniker in table.EnumRunning():
            try:
                name = moniker.GetDisplayName(context, None).lstrip('!')
                if path_key(name) != wanted:
                    continue
                book = win32com.client.Dispatch(table.GetObject(moniker))
                if path_key(book.FullName) == wanted:
                    return book.Application, book
            except Exception:
                continue
    except Exception:
        pass
    return None, None


class SourceWorkbookWriteback:
    def __init__(self, engine, source, baseline, sheet_name):
        self.source = Path(source)
        self.sheet_name = sheet_name
        self.headers = list(engine.STATUS_HEADERS)
        self.normalize = engine.normalize_header
        self.pending = {}
        self.written = set()
        self.attempted = {}
        self.message = ''
        self.app = self.book = None
        self.owned = False
        book = load_workbook(baseline, read_only=True, data_only=False)
        try:
            self.epoch = book.epoch
            sheet = book[sheet_name]
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            book.close()
        self.baseline_headers = list(rows[0])
        self.baseline_rows = {number: list(values) for number, values in enumerate(rows[1:], 2)}
        self.status_keys = {self.normalize(header) for header in self.headers}
        self.input_columns = {index: self.normalize(header) for index, header in enumerate(self.baseline_headers)
                              if self.normalize(header) and self.normalize(header) not in self.status_keys}
        if len(set(self.input_columns.values())) != len(self.input_columns):
            raise ValueError('Original workbook contains duplicate input headers.')

    def record(self, store, rows, updates):
        for row in rows:
            changed = {header: store.worksheet.cell(row, store.status_columns[header]).value
                       for header, value in updates.items() if value is not None and header in self.headers}
            if changed:
                changed['Last Run Time'] = store.worksheet.cell(row, store.status_columns['Last Run Time']).value
                self.pending.setdefault(row, {}).update(changed)

    def summary(self):
        return dict(sourcePath=str(self.source), status='pending' if self.pending else 'synced' if self.written else 'unchanged',
                    writtenRows=len(self.written), pendingRows=len(self.pending), message=self.message)

    def connect(self):
        if self.book is not None:
            # A stale/disconnected workbook is not proof that the original is
            # available for rewriting. The caller retains pending results.
            if path_key(self.book.FullName) != path_key(self.source):
                raise ValueError('The Excel workbook identity changed.')
            return
        if self.app is None:
            self.app, self.book = attach_workbook(self.source)
            if self.book is not None:
                return
        import win32com.client
        if self.app is None:
            self.app = win32com.client.DispatchEx('Excel.Application')
            self.owned = True
        # These settings affect only our private instance. Disable macros and
        # link updates before opening, not user application-wide preferences.
        self.app.AutomationSecurity = 3
        self.app.DisplayAlerts = False
        self.app.EnableEvents = False
        self.book = self.app.Workbooks.Open(Filename=str(self.source.resolve()), UpdateLinks=0,
                                           ReadOnly=False, Notify=False, AddToMru=False,
                                           IgnoreReadOnlyRecommended=True)

    def flush(self):
        if not self.pending:
            return self.summary()
        try:
            self.connect()
            if path_key(self.book.FullName) != path_key(self.source):
                raise ValueError('The Excel workbook identity changed.')
            if self.book.ReadOnly:
                raise ValueError('Original workbook is read-only or locked by another Excel instance.')
            if not self.app.Ready:
                raise ValueError('Excel is busy. Finish editing the current cell or dialog.')
            if bool(self.book.Date1904) != (self.epoch.year == 1904):
                raise ValueError('The workbook date system changed.')
            sheet = self.book.Worksheets(self.sheet_name)
            if sheet.ProtectContents:
                raise ValueError('The original worksheet is protected.')
            count = max(1, int(sheet.UsedRange.Column) + int(sheet.UsedRange.Columns.Count) - 1)
            if count > 16384:
                raise ValueError('Invalid Excel column count.')
            header_values = row_values(sheet, 1, count)
            columns = {}
            for column, header in enumerate(header_values, 1):
                key = self.normalize(header)
                if key:
                    if key in columns:
                        raise ValueError('Original workbook contains duplicate headers.')
                    columns[key] = column
            if any(key not in columns for key in self.input_columns.values()):
                raise ValueError('Original input headers changed. Results were not written to a different layout.')
            for header in self.headers:
                key = self.normalize(header)
                if key not in columns:
                    count += 1
                    columns[key] = count

            changes = []
            # Prevalidate every affected row before touching any cell. A sort,
            # changed Material/Parma/Qty/etc., or new RFQ number cannot silently
            # receive another record's results. Formula text is checked too.
            current_count = len(header_values)
            for row, updates in sorted(self.pending.items()):
                baseline = self.baseline_rows.get(row)
                if baseline is None:
                    raise ValueError(f'Excel row {row} was not in the validated source.')
                current = row_values(sheet, row, current_count)
                current.extend([None] * (count - current_count))
                for original_index, key in self.input_columns.items():
                    if comparable(current[columns[key] - 1], self.epoch) != comparable(baseline[original_index], self.epoch):
                        raise ValueError(f'Excel row {row} input changed or moved ({self.baseline_headers[original_index]}); no result was written to that row.')
                baseline_status = {self.normalize(header): baseline[index] for index, header in enumerate(self.baseline_headers)
                                   if self.normalize(header) in self.status_keys}
                for header, value in updates.items():
                    column = columns[self.normalize(header)]
                    old = comparable(current[column - 1], self.epoch)
                    allowed = [comparable(baseline_status.get(self.normalize(header)), self.epoch),
                               comparable(value, self.epoch)]
                    if (row, header) in self.attempted:
                        allowed.append(comparable(self.attempted[row, header], self.epoch))
                    if old not in allowed or isinstance(current[column - 1], ExcelFormula):
                        raise ValueError(f'Excel row {row} result column {header} changed; existing results were preserved.')
                    changes.append((row, column, header, value))

            for header in self.headers:
                column = columns[self.normalize(header)]
                if column > current_count:
                    sheet.Cells(1, column).Value2 = header
            for row, column, header, value in changes:
                # Remember attempted values before dispatch: an Excel write or
                # Save may succeed even when COM subsequently reports an error.
                self.attempted[row, header] = value
                cell = sheet.Cells(row, column)
                if isinstance(value, str):
                    cell.NumberFormat = '@'  # RFQ IDs keep zeros; text is not a formula.
                cell.Value2 = value
            self.book.Save()
            if not self.book.Saved:
                raise ValueError('Excel has not confirmed the original workbook was saved.')
            for row, column, _, value in changes:
                if comparable(sheet.Cells(row, column).Value2, self.epoch) != comparable(value, self.epoch):
                    raise ValueError(f'Excel row {row} result read-back did not match.')
            self.written.update(self.pending)
            self.pending.clear()
            self.message = 'RFQ results were saved to the original Excel workbook.'
        except Exception as exc:
            self.message = ('Original Excel could not be updated. Results are safe in the result copy; '
                            'do not rerun SAP to retry Excel writeback. ' + str(exc))
        return self.summary()

    def close(self):
        try:
            return self.flush()
        finally:
            if self.owned:
                try:
                    # A background workbook can become user-visible during a
                    # long run. Never close that window or discard unsaved edits,
                    # even though we originally launched its Excel process.
                    safe_to_close = not self.app.Visible and (self.book is None or self.book.Saved)
                    if self.book is not None and safe_to_close:
                        self.book.Close(SaveChanges=False)
                    elif self.book is not None:
                        self.app.Visible = True
                    if safe_to_close and int(self.app.Workbooks.Count) == 0:
                        self.app.Quit()
                except Exception:
                    pass
            self.app = self.book = None
